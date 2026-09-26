"""Build L3M1 Saplings Lessons 18-22 as RemoEd PPTX decks.

Same engine as build_l3m1_lessons_13_15.py (itself generalised from build_l3m1_lesson_12.py), plus optional
per-lesson "title_lines" for long titles. Chrome matches the laptop's shared composer
(build_l2m1_lessons_7_9.compose_slide): white lesson pill (0.25",0.2") 4.2"x0.45" #1A5696,
indigo #5A67D8 page badge (11.55",0.2") 1.5"x0.45", yellow #FAD648 learner chip (#1E3A5F),
white footer label + knocked-out logo with drop shadows and NO card, 1920x1080 composed JPG
placed full-bleed on a 13.333"x7.5" slide. Deliberate differences (same as Lesson 12):
  * title slide: ~52-58pt title as two stacked white pills low-left (not a mid-frame banner
    that covers faces); greeting chip in the open top band;
  * wide body chips are lifted above the footer row instead of colliding with it;
  * art is centre-cropped to 16:9 before resizing (never stretched);
  * phonics / card words are stamped in code (README: never bake text into art);
  * nothing is ever overwritten without a timestamped backup in a sibling "_backup" folder,
    and the D: Level folder is never created (copied only if it exists).

Art: docs/lesson-references/lessons/L3M1-Lesson-<N>/art/l3m1-l<N>-slide-01.png ... -18.png
Stamp positions can be tuned without editing this file: put a stamps.json next to the art
folder ({"4": [{"kind": "lines", ...}], ...}); it replaces that lesson's STAMPS per slide.

Usage (laptop, from docs/lesson-references):
    python build_l3m1_lessons_18_22.py --lessons 18 19 20     (batch 1)
    python build_l3m1_lessons_18_22.py --lessons 21 22        (batch 2)
    python build_l3m1_lessons_18_22.py --lessons 18 --stamp-debug --no-copies   (tuning pass)
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Emu, Inches

FOOTER = "SAPLINGS! MONTH 1"
LEVEL_NUM = 3
DEFAULT_ROOT = Path(r"D:\Users\Window11\Documents\RemoEd-Jean\RemoEdPH")
DEFAULT_LAPTOP = Path(r"D:\Users\Window11\Desktop\JeanDesktop\RemoEdPH\A Lesson and Training Materials")

NAVY = (26, 86, 150)
NAVY_DARK = (30, 58, 95)
INDIGO = (90, 103, 216)
YELLOW = (250, 214, 72)
WHITE = (255, 255, 255)
PLAQUE_INK = (255, 240, 205)
PLAQUE_HALO = (45, 25, 10)

# (on-screen text, is_title, chip position) - chip: "bottom" | "top" | (centre-x frac, top-y frac)
# On-screen text is verbatim from the brief; quotation marks dropped to match laptop L3M1 Lessons 1-6.
# "title_lines" (optional) sets how the title-slide pills break, for titles too long for "Lesson N -" / "Title".
# Stamp boxes for Lessons 18-20 are measured on the final art; 21-22 are still provisional (re-measure with --stamp-debug).


def _row(words, centres, y0, y1, hw):
    return [{"kind": "word", "text": w, "box": (c - hw, y0, c + hw, y1), "group": "row"} for w, c in zip(words, centres)]


LESSONS = {
    18: {
        "title": "Lesson 18 \u2013 Review: Self-Worth & Respect",
        "title_lines": ["Lesson 18 \u2013", "Review: Self-Worth", "& Respect"],   # 3 short pills stay left of Remo
        "title_bottom": 6.9,
        "title_max_right": 0.615,   # final art: Remo's face sits at x 0.64-0.72, y 0.62-0.80 on slide 1
        "folder": "L3M1-Lesson-18",
        "pptx": "L3M1-Lesson-18-Review-Self-Worth-and-Respect.pptx",
        "prefix": "l3m1-l18-slide-",
        "pages": [
            ("Worth & Respect Celebration!", True, (0.5, 0.10)),
            ("Crown & Heart!", False, "bottom"),
            ("Unique & Special", False, "bottom"),
            ("Respect & Kind Words", False, "bottom"),
            ("Kindness Wins!", False, "bottom"),
            ("God made me special, and I respect others.", False, "bottom"),
            ("I am unique because...", False, "bottom"),
            ("Gentle Hands, Fair Turns!", False, "bottom"),
            ("Say Something Kind!", False, "bottom"),
            ("Everyone is valuable to God.", False, "bottom"),
            ("Which picture shows respect?", False, "bottom"),
            ("You Are Valuable!", False, "bottom"),
            ("Catch the Kindness Stars!", False, "bottom"),
            ("Unique. Special. Respect. Kind.", False, (0.42, 0.81)),   # below the cards (they end at y 0.78)
            ("Worth & Respect Mastered!", False, "bottom"),
            ("God made me special, and I respect others.", False, "bottom"),
            ("Golden Crown Badge Unlocked!", False, "bottom"),
            ("Goodbye!", False, "bottom"),
        ],
        "stamps": {
            # measured on final art: sky/roof corner of each panel; four cards x 0.13-0.87, y 0.31-0.78
            11: [{"kind": "badge", "text": "A", "center": (0.055, 0.175), "r": 0.036},
                 {"kind": "badge", "text": "B", "center": (0.555, 0.175), "r": 0.036}],
            14: _row(["Unique", "Special", "Respect", "Kind"], [0.215, 0.4075, 0.5975, 0.785], 0.47, 0.62, 0.07),
        },
    },
    19: {
        "title": "Lesson 19 \u2013 Review: Home & Family Nouns",
        "title_lines": ["Lesson 19 \u2013 Review:", "Home & Family Nouns"],
        "title_bottom": 6.75,   # final art: Ed's face ends at y 0.56 on slide 1, so the pills sit a little lower
        "folder": "L3M1-Lesson-19",
        "pptx": "L3M1-Lesson-19-Review-Home-and-Family-Nouns.pptx",
        "prefix": "l3m1-l19-slide-",
        "pages": [
            ("Home & Family Review!", True, (0.5, 0.10)),
            ("Roof & Hug!", False, "bottom"),
            ("Bedroom", False, "bottom"),
            ("Kitchen", False, "bottom"),
            ("Living Room", False, "bottom"),
            ("Parents", False, "bottom"),
            ("Siblings & Grandparents", False, "bottom"),
            ("I love my family and my home.", False, "bottom"),
            ("Where is the cooking done?", False, "bottom"),
            ("Thank God for home and family.", False, "bottom"),
            ("Complete the Family Tree!", False, "bottom"),
            ("What room are you in?", False, "bottom"),
            ("Home & Family Expert!", False, "bottom"),
            ("Speed Nouns!", False, (0.5, 0.83)),   # below the cards (they end at y 0.81)
            ("Nouns Mastered!", False, "bottom"),
            ("I love my family and my home.", False, "bottom"),
            ("Home & Family Badge Unlocked!", False, "bottom"),
            ("Goodbye!", False, "bottom"),
        ],
        "stamps": {
            # measured on final art; dark-wood plaques get cream lettering with a dark halo
            3: [{"kind": "word", "text": "Bedroom", "box": (0.405, 0.275, 0.50, 0.365), "color": PLAQUE_INK, "halo": PLAQUE_HALO}],
            4: [{"kind": "word", "text": "Kitchen", "box": (0.30, 0.35, 0.51, 0.47), "color": PLAQUE_INK, "halo": PLAQUE_HALO}],
            5: [{"kind": "word", "text": "Living\nRoom", "box": (0.435, 0.195, 0.555, 0.305), "color": PLAQUE_INK, "halo": PLAQUE_HALO}],
            # white label strips y 0.60-0.665 under Ed / toddler sister / grandmother / grandfather
            7: [{"kind": "word", "text": w, "box": (a + 0.006, 0.607, b - 0.006, 0.658), "group": "row"} for w, (a, b) in
                zip(["Brother", "Sister", "Grandmother", "Grandfather"], [(0.11, 0.24), (0.33, 0.46), (0.545, 0.675), (0.765, 0.895)])],
            9: [{"kind": "badge", "text": "A", "center": (0.055, 0.175), "r": 0.036},
                {"kind": "badge", "text": "B", "center": (0.555, 0.175), "r": 0.036}],
            # five cards y 0.50-0.81; words sit in the upper-middle so the chip/footer never touch them
            14: [{"kind": "word", "text": w, "box": (a + 0.012, 0.56, b - 0.012, 0.69), "group": "row"} for w, (a, b) in
                 zip(["House", "Bedroom", "Kitchen", "Parents", "Family"],
                     [(0.12, 0.26), (0.28, 0.415), (0.44, 0.575), (0.595, 0.73), (0.755, 0.89)])],
        },
    },
    20: {
        # target words confirmed by Kristine: the p.106 slide 14 set (Name, Cat, Bag, Help, Family), not the p.104 header list
        "title": "Lesson 20 \u2013 Integrated Story Challenge",
        "folder": "L3M1-Lesson-20",
        "pptx": "L3M1-Lesson-20-Integrated-Story-Challenge.pptx",
        "prefix": "l3m1-l20-slide-",
        "pages": [
            ("Grand Story Challenge!", True, (0.5, 0.10)),
            ("Pointer Finger Up!", False, "bottom"),
            ("My name is Sam. God made me unique!", False, "bottom"),
            ("The cat is on the mat. C-A-T, cat.", False, (0.40, 0.84)),   # low-left: clears the cat's face (x 0.66-0.82, y 0.60-0.80)
            ("I have a red hat in my bag.", False, "bottom"),
            ("May I use the device, please?", False, "bottom"),
            ("Screen time is finished.", False, "bottom"),
            ("I help my family in the kitchen.", False, "bottom"),
            ("I love my family. We are all different and good.", False, "bottom"),
            ("God helped me learn and grow!", False, "bottom"),
            ("What was on the mat?", False, (0.5, 0.10)),   # top: the Cat / Bag plates use the bottom band
            ("Screen time is ____.", False, "bottom"),
            ("Story Challenge Passed!", False, "bottom"),
            ("Month 1 Word Cloud", False, "bottom"),
            ("Month 1 Story Master!", False, "bottom"),
            ("My name is Sam. I love my family.", False, "bottom"),
            ("Grand Story Badge Unlocked!", False, "bottom"),
            ("Goodbye!", False, "bottom"),
        ],
        "stamps": {
            3: [{"kind": "word", "text": "Sam", "box": (0.437, 0.422, 0.495, 0.472)}],     # Sam's name tag (measured)
            # no strips in the art: white rounded label plates drawn in code on the front of each mat
            11: [{"kind": "word", "text": "Cat", "box": (0.15, 0.80, 0.35, 0.88), "plate": True, "group": "row"},
                 {"kind": "word", "text": "Bag", "box": (0.65, 0.80, 0.85, 0.88), "plate": True, "group": "row"}],
            14: [{"kind": "word", "text": w, "box": b, "color": c} for w, b, c in [          # word cloud (board x 0.26-0.76, y 0.16-0.77)
                ("Name", (0.35, 0.29, 0.50, 0.40), (30, 58, 95)),
                ("Cat", (0.525, 0.28, 0.655, 0.39), (214, 96, 30)),
                ("Family", (0.30, 0.435, 0.52, 0.55), (90, 103, 216)),
                ("Bag", (0.555, 0.435, 0.715, 0.55), (200, 50, 60)),
                ("Help", (0.42, 0.575, 0.585, 0.685), (40, 140, 70))]],
        },
    },
    21: {
        "title": "Lesson 21 \u2013 App Garden Milestone 1",
        "folder": "L3M1-Lesson-21",
        "pptx": "L3M1-Lesson-21-App-Garden-Milestone-1.pptx",
        "prefix": "l3m1-l21-slide-",
        "pages": [
            ("Welcome to Garden Milestone 1!", True, (0.5, 0.10)),
            ("Raindrops & Sunbeams!", False, "bottom"),
            ("Garden", False, "bottom"),
            ("Grow", False, "bottom"),
            ("Water", False, "bottom"),
            ("My Eco-Drops make my garden grow.", False, "bottom"),
            ("Eco-Drop Vault: 20 Drops Ready!", False, "bottom"),
            ("Tap to Water Your Garden!", False, "bottom"),
            ("Month 1 Flower Blooms!", False, "bottom"),
            ("Good learning brings good fruit.", False, "bottom"),
            ("Choose Garden Decorations!", False, "bottom"),
            ("Garden Snapshot Saved!", False, "bottom"),
            ("Show Mom & Dad Your Garden!", False, "bottom"),
            ("Garden. Grow. Plant. Water.", False, "bottom"),
            ("Milestone 1 Complete!", False, "bottom"),
            ("My Eco-Drops make my garden grow.", False, "bottom"),
            ("Master Gardener Badge Unlocked!", False, "bottom"),
            ("Goodbye!", False, "bottom"),
        ],
        "stamps": {
            14: _row(["Garden", "Grow", "Plant", "Water"], [0.20, 0.40, 0.60, 0.80], 0.36, 0.54, 0.075),
        },
    },
    22: {
        "title": "Lesson 22 \u2013 Month 1 Grand Showcase",
        "folder": "L3M1-Lesson-22",
        "pptx": "L3M1-Lesson-22-Month-1-Grand-Showcase.pptx",
        "prefix": "l3m1-l22-slide-",
        "pages": [
            ("Month 1 Grand Showcase!", True, (0.5, 0.10)),
            ("Strike Your Superstar Pose!", False, "bottom"),
            ("Task 1: Write Your First Name!", False, "bottom"),
            ("Writing Name...", False, "bottom"),
            ("Show Your Written Name!", False, "bottom"),
            ("Task 2: Read the Showcase Sentence!", False, "bottom"),
            ("The cat sat on the mat with a hat.", False, "bottom"),
            ("Task 3: State Your Value & Tech Rule!", False, "bottom"),
            ("All glory to God for our growth!", False, "bottom"),
            ("Evaluating Showcase... PASSED!", False, "bottom"),
            ("Month 1 Badge Awarded!", False, "bottom"),
            ("Parent High-Five Ceremony!", False, "bottom"),
            ("My name is [Name], and I can read!", False, "bottom"),
            ("Month 1 Complete!", False, "bottom"),
            ("Get Ready for Month 2!", False, "bottom"),
            ("Congratulations Saplings!", False, "bottom"),
            ("Month 1 Master Trophy Saved!", False, "bottom"),
            ("Goodbye & See You in Month 2!", False, "bottom"),
        ],
        "stamps": {
            3: [{"kind": "dotted", "text": "Name", "box": (0.30, 0.22, 0.70, 0.52)}],       # task whiteboard
            5: [{"kind": "word", "text": "Ed", "box": (0.40, 0.47, 0.60, 0.65)}],           # Ed's written paper
            7: [{"kind": "word", "text": "The cat sat on the mat with a hat.", "box": (0.20, 0.60, 0.80, 0.72)}],   # cream strip
            14: [{"kind": "word", "text": w, "box": (0.35, c - 0.04, 0.74, c + 0.04), "align": "left", "group": "rows"}   # checklist rows
                 for w, c in zip(["Name", "Phonics", "Tech Stewardship", "Values", "Home Nouns"], [0.22, 0.33, 0.44, 0.55, 0.66])],
        },
    },
}

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
PX_W, PX_H = 1920, 1080


def ix(inches: float) -> int:
    return int(round(inches * PX_W / 13.333))


def iy(inches: float) -> int:
    return int(round(inches * PX_H / 7.5))


class Ctx:
    font_path: str | None = None
    fallback_path: str | None = None
    logo: Path | None = None
    stamp_debug = False
    stamp_ts = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")


def first_existing(paths):
    for p in paths:
        if p and Path(p).exists():
            return str(p)
    return None


def load_font(size: int, fallback: bool = False):
    path = Ctx.fallback_path if fallback else Ctx.font_path
    return ImageFont.truetype(path, size) if path else ImageFont.load_default()


_MISSING: dict = {}


def has_glyph(font, ch: str) -> bool:
    key = (id(font), ch)
    if key not in _MISSING:
        notdef = font.getmask("\ue000")
        m = font.getmask(ch)
        _MISSING[key] = not (m.size == notdef.size and bytes(m) == bytes(notdef))
    return _MISSING[key]


def runs(text: str, size: int):
    """Split text into (segment, font) runs, using the fallback font for glyphs the main font lacks (e.g. U+2192)."""
    main = load_font(size)
    out, cur, cur_fb = [], "", False
    for ch in text:
        fb = bool(Ctx.fallback_path) and not has_glyph(main, ch)
        if cur and fb != cur_fb:
            out.append((cur, load_font(size, cur_fb)))
            cur = ""
        cur += ch
        cur_fb = fb
    if cur:
        out.append((cur, load_font(size, cur_fb)))
    return out


def measure(draw, text: str, size: int):
    w, top, bot = 0, 10 ** 6, -10 ** 6
    for seg, f in runs(text, size):
        b = draw.textbbox((w, 0), seg, font=f)
        w = b[2]
        top, bot = min(top, b[1]), max(bot, b[3])
    return w, bot - top, top


def draw_runs(draw, xy, text: str, size: int, fill):
    x, y = xy
    for seg, f in runs(text, size):
        draw.text((x, y), seg, font=f, fill=fill)
        x = draw.textbbox((x, y), seg, font=f)[2]


def text_size(draw, text, font):
    b = draw.textbbox((0, 0), text, font=font)
    return b[2] - b[0], b[3] - b[1]


def round_rect(draw, xy, radius, fill):
    draw.rounded_rectangle(xy, radius=radius, fill=fill)


# ---------------------------------------------------------------- safe file writes
def backup_if_exists(dest: Path):
    if dest.exists():
        bdir = dest.parent / "_backup"
        bdir.mkdir(exist_ok=True)
        b = bdir / f"{dest.stem}.{Ctx.stamp_ts}{dest.suffix}"
        shutil.copy2(dest, b)
        print("backup", b)


LOCKED: list = []


def safe_copy(src: Path, dest: Path) -> bool:
    """Copy with backup; a file locked by PowerPoint etc. is reported and skipped, not fatal."""
    try:
        backup_if_exists(dest)
        shutil.copy2(src, dest)
        return True
    except PermissionError as e:
        print("LOCKED (not updated):", dest, "-", e)
        LOCKED.append(str(dest))
        return False


# ---------------------------------------------------------------- footer / chrome
def knock_out_white(im):
    im = im.convert("RGBA")
    px = im.load()
    w, h = im.size
    for y in range(h):
        for x in range(w):
            r, g, b, a = px[x, y]
            if r > 246 and g > 246 and b > 246:
                px[x, y] = (r, g, b, 0)
            elif r > 228 and g > 228 and b > 228:
                px[x, y] = (r, g, b, max(0, a - int((min(r, g, b) - 228) / 18 * 255)))
    return im


_LOGO: dict = {}


def footer_geometry(draw):
    tw, th = text_size(draw, FOOTER, load_font(34))
    logo_w, logo_h = ix(0.78), iy(0.78)
    logo_x = PX_W - ix(0.28) - logo_w
    logo_y = PX_H - iy(0.22) - logo_h
    text_x = logo_x - ix(0.14) - tw
    return tw, th, logo_w, logo_h, logo_x, logo_y, text_x


def draw_footer(img):
    canvas = img.convert("RGBA")
    tw, th, logo_w, logo_h, logo_x, logo_y, text_x = footer_geometry(ImageDraw.Draw(canvas))
    text_y = logo_y + (logo_h - th) // 2 - 2
    font = load_font(34)
    shadow = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    ImageDraw.Draw(shadow).text((text_x, text_y + 2), FOOTER, font=font, fill=(0, 0, 0, 179))  # rgba(0,0,0,0.7)
    canvas = Image.alpha_composite(canvas, shadow.filter(ImageFilter.GaussianBlur(4)))
    sharp = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    ImageDraw.Draw(sharp).text((text_x, text_y), FOOTER, font=font, fill=WHITE + (255,))
    canvas = Image.alpha_composite(canvas, sharp)
    if "logo" not in _LOGO:
        _LOGO["logo"] = knock_out_white(Image.open(Ctx.logo)).resize((logo_w, logo_h), Image.Resampling.LANCZOS)
    logo = _LOGO["logo"]
    ls = Image.new("RGBA", logo.size, (0, 0, 0, 0))
    ls.putalpha(logo.split()[-1].point(lambda a: int(a * 0.6)))  # rgba(0,0,0,0.6)
    pad = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    pad.paste(ls.filter(ImageFilter.GaussianBlur(3)), (logo_x, logo_y + 2), ls.filter(ImageFilter.GaussianBlur(3)))
    canvas = Image.alpha_composite(canvas, pad)
    canvas.paste(logo, (logo_x, logo_y), logo)
    return canvas.convert("RGB")


# ---------------------------------------------------------------- code-stamped text
def fit_size(draw, lines, w, h, start=400, floor=18):
    size = start
    while size > floor:
        f = load_font(size)
        lw = max(text_size(draw, ln, f)[0] for ln in lines)
        lh = int(size * 1.08) * len(lines)
        if lw <= w and lh <= h:
            return size
        size -= 4
    return floor


def stamp(img, s):
    draw = ImageDraw.Draw(img)
    kind = s["kind"]
    if kind == "badge":
        cx, cy = int(PX_W * s["center"][0]), int(PX_H * s["center"][1])
        r = int(PX_W * s.get("r", 0.04))
        sh = Image.new("RGBA", img.size, (0, 0, 0, 0))
        ImageDraw.Draw(sh).ellipse((cx - r, cy - r + 4, cx + r, cy + r + 4), fill=(0, 0, 0, 110))
        img = Image.alpha_composite(img.convert("RGBA"), sh.filter(ImageFilter.GaussianBlur(5))).convert("RGB")
        draw = ImageDraw.Draw(img)
        draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill=YELLOW, outline=WHITE, width=6)
        f = load_font(int(r * 1.2))
        b = draw.textbbox((0, 0), s["text"], font=f)
        draw.text((cx - (b[0] + b[2]) // 2, cy - (b[1] + b[3]) // 2), s["text"], font=f, fill=NAVY_DARK)
        return img
    x0, y0, x1, y1 = (int(PX_W * s["box"][0]), int(PX_H * s["box"][1]), int(PX_W * s["box"][2]), int(PX_H * s["box"][3]))
    if Ctx.stamp_debug:
        draw.rectangle((x0, y0, x1, y1), outline=(255, 0, 0), width=3)
    if kind == "lines":
        n = s.get("n", 3)
        gap = (x1 - x0) * 0.12
        seg = ((x1 - x0) - gap * (n - 1)) / n
        th = max(12, (y1 - y0) // 3)
        for k in range(n):
            lx = int(x0 + k * (seg + gap))
            round_rect(draw, (lx - 2, y1 - th - 2, int(lx + seg) + 2, y1 + 2), th // 2 + 2, WHITE)
            round_rect(draw, (lx, y1 - th, int(lx + seg), y1), th // 2, NAVY_DARK)
        return img
    if s.get("plate"):   # clean white rounded label strip drawn in code (art has none)
        rad = (y1 - y0) // 3
        sh = Image.new("RGBA", img.size, (0, 0, 0, 0))
        round_rect(ImageDraw.Draw(sh), (x0, y0 + 6, x1, y1 + 6), rad, (0, 0, 0, 90))
        img = Image.alpha_composite(img.convert("RGBA"), sh.filter(ImageFilter.GaussianBlur(6))).convert("RGB")
        draw = ImageDraw.Draw(img)
        round_rect(draw, (x0, y0, x1, y1), rad, WHITE)
        pad_x, pad_y = int((x1 - x0) * 0.12), int((y1 - y0) * 0.16)
        x0, y0, x1, y1 = x0 + pad_x, y0 + pad_y, x1 - pad_x, y1 - pad_y
    lines = s["text"].split("\n")
    size = s.get("size") or fit_size(draw, lines, x1 - x0, y1 - y0)
    f = load_font(size)
    lh = int(size * 1.08)
    total_h = lh * len(lines)
    layer = Image.new("L", img.size, 0)
    ld = ImageDraw.Draw(layer)
    for k, ln in enumerate(lines):
        b = ld.textbbox((0, 0), ln, font=f)
        tx = (x0 - b[0]) if s.get("align") == "left" else x0 + ((x1 - x0) - (b[2] - b[0])) // 2 - b[0]
        ty = y0 + ((y1 - y0) - total_h) // 2 + k * lh - b[1] + (lh - (b[3] - b[1])) // 2
        ld.text((tx, ty), ln, font=f, fill=255)
    base = img.convert("RGBA")
    if kind == "glow":
        glow = Image.new("RGBA", img.size, (255, 236, 140, 0))
        glow.putalpha(layer.filter(ImageFilter.MaxFilter(9)).filter(ImageFilter.GaussianBlur(14)))
        base = Image.alpha_composite(base, glow)
        base = Image.alpha_composite(base, glow)
        core = Image.new("RGBA", img.size, (255, 250, 220, 0))
        core.putalpha(layer)
        base = Image.alpha_composite(base, core)
    elif kind == "dotted":
        # tracing style: faint letter fill + dashed outline ring
        from PIL import ImageChops
        k = max(7, (size // 18) | 1)
        ring = ImageChops.subtract(layer.filter(ImageFilter.MaxFilter(k)), layer.filter(ImageFilter.MinFilter(k)))
        step = max(12, size // 9)
        dots = Image.new("L", img.size, 0)
        dd = ImageDraw.Draw(dots)
        for yy in range(y0 - step, y1 + step, step):
            for xx in range(x0 - step, x1 + step, step):
                dd.ellipse((xx, yy, xx + step * 0.62, yy + step * 0.62), fill=255)
        fill = Image.new("RGBA", img.size, (160, 170, 240, 0))
        fill.putalpha(layer.point(lambda v: int(v * 0.35)))
        base = Image.alpha_composite(base, fill)
        col = Image.new("RGBA", img.size, INDIGO + (0,))
        col.putalpha(ImageChops.multiply(ring, dots))
        base = Image.alpha_composite(base, col)
    else:  # word: navy text with a soft white halo so it reads on boards, cards and paper
        halo = Image.new("RGBA", img.size, tuple(s.get("halo", WHITE)) + (0,))
        halo.putalpha(layer.filter(ImageFilter.MaxFilter(5)).point(lambda a: int(a * 0.55)))
        base = Image.alpha_composite(base, halo)
        txt = Image.new("RGBA", img.size, tuple(s.get("color", NAVY_DARK)) + (0,))
        txt.putalpha(layer)
        base = Image.alpha_composite(base, txt)
    return base.convert("RGB")


def uniform_sizes(img, stamps):
    """Stamps sharing a "group" key get the same (largest common) font size, so a row of cards reads evenly."""
    groups: dict = {}
    draw = ImageDraw.Draw(img)
    for st in stamps:
        if st.get("group") and "box" in st and not st.get("size"):
            x0, y0, x1, y1 = (int(PX_W * st["box"][0]), int(PX_H * st["box"][1]), int(PX_W * st["box"][2]), int(PX_H * st["box"][3]))
            sz = fit_size(draw, st["text"].split("\n"), x1 - x0, y1 - y0)
            groups[st["group"]] = min(groups.get(st["group"], 10 ** 6), sz)
    return [dict(st, size=groups[st["group"]]) if st.get("group") in groups else st for st in stamps]


# ---------------------------------------------------------------- slide composition
def load_art(src: Path):
    img = Image.open(src).convert("RGB")
    w, h = img.size
    t = PX_W / PX_H
    if abs(w / h - t) > 0.01:
        if w / h > t:
            nw = int(h * t)
            img = img.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
        else:
            nh = int(w / t)
            img = img.crop((0, (h - nh) // 2, w, (h - nh) // 2 + nh))
    return img.resize((PX_W, PX_H), Image.Resampling.LANCZOS)


def compose_slide(src, dest, title, page, total, phrase, is_title, chip_pos, stamps, title_bottom=6.42, title_lines=None,
                  title_max_right=None):
    img = load_art(src)
    stamps = uniform_sizes(img, stamps)
    for s in stamps:
        img = stamp(img, s)
    draw = ImageDraw.Draw(img)
    font14 = load_font(36)

    tw, th = text_size(draw, title, font14)
    lw = max(ix(4.2), tw + 44)
    lx, ly = ix(0.25), iy(0.2)
    round_rect(draw, (lx, ly, lx + lw, ly + iy(0.45)), 22, WHITE)
    draw.text((lx + 18, ly + (iy(0.45) - th) // 2 - 2), title, font=font14, fill=NAVY)

    badge = f"{page} / {total}"
    bw, bh, bx, by = ix(1.5), iy(0.45), ix(11.55), iy(0.2)
    round_rect(draw, (bx, by, bx + bw, by + bh), 22, INDIGO)
    tw, th = text_size(draw, badge, font14)
    draw.text((bx + (bw - tw) // 2, by + (bh - th) // 2 - 2), badge, font=font14, fill=WHITE)

    img = draw_footer(img)
    draw = ImageDraw.Draw(img)
    f_tw, _, _, _, _, logo_y, footer_left = footer_geometry(draw)

    if is_title:
        # two stacked pills low-left: "Lesson N -" / "{Title}", 52-58pt (104-116px)
        lines = list(title_lines) if title_lines else [t.strip() for t in title.replace("\u2013", "\u2013\n", 1).split("\n")]
        max_right = footer_left - ix(0.2)
        if title_max_right:   # keep the pills left of a face that sits low in the frame (e.g. Remo)
            max_right = min(max_right, int(PX_W * title_max_right))
        size = 116
        while True:
            f = load_font(size)
            widths = [draw.textbbox((0, 0), ln, font=f)[2] for ln in lines]
            if ix(0.3) + max(widths) + 70 <= max_right or size <= 104:
                break
            size -= 4
        asc, desc = f.getmetrics()
        pill_h, gap = asc + desc + 24, 12
        bottom = iy(title_bottom)
        right = ix(0.3) + max(widths) + 70
        if bottom > logo_y - 12 and right > footer_left - ix(0.15):
            bottom = min(bottom, logo_y - 12)   # never let a low title collide with the footer
        y = bottom - len(lines) * pill_h - (len(lines) - 1) * gap
        for ln, w in zip(lines, widths):
            round_rect(draw, (ix(0.3), y, ix(0.3) + w + 70, y + pill_h), 28, WHITE)
            draw.text((ix(0.3) + 35, y + 12), ln, font=f, fill=NAVY)
            y += pill_h + gap

    size = 64
    tw, th, top = measure(draw, phrase, size)
    cw, ch = tw + 72, max(iy(0.7), th + 32)
    if isinstance(chip_pos, tuple):
        cx, cy = int(PX_W * chip_pos[0]) - cw // 2, int(PX_H * chip_pos[1])
    elif chip_pos == "top":
        cx, cy = (PX_W - cw) // 2, iy(0.9)
    else:
        cx, cy = (PX_W - cw) // 2, iy(6.05)
        if cx + cw > footer_left - ix(0.15):
            cy = min(cy, logo_y - 18 - ch)
    round_rect(draw, (cx, cy, cx + cw, cy + ch), 26, YELLOW)
    draw_runs(draw, (cx + (cw - tw) // 2, cy + (ch - th) // 2 - top), phrase, size, NAVY_DARK)

    dest.parent.mkdir(parents=True, exist_ok=True)
    backup_if_exists(dest)
    img.save(dest, "JPEG", quality=92)


def build(n: int, refs: Path, laptop_root: Path, copies: bool, only=None):
    L = LESSONS[n]
    folder = refs / "lessons" / L["folder"]
    art_dir, comp_dir = folder / "art", folder / "composed"
    total = len(L["pages"])
    stamps = {int(k): v for k, v in L.get("stamps", {}).items()}
    override = folder / "stamps.json"
    if override.exists():
        for k, v in json.loads(override.read_text(encoding="utf-8")).items():
            stamps[int(k)] = [dict(s, box=tuple(s["box"])) if "box" in s else dict(s, center=tuple(s["center"])) for s in v]
        print("stamp overrides from", override)
    missing = [i for i in range(1, total + 1) if not (art_dir / f"{L['prefix']}{i:02d}.png").exists()]
    if missing:
        raise SystemExit(f"Lesson {n}: missing art for slides {missing} in {art_dir}")
    composed = []
    for i, (phrase, is_title, pos) in enumerate(L["pages"], 1):
        dest = comp_dir / f"{L['prefix']}{i:02d}.jpg"
        if only is None or i in only:
            compose_slide(art_dir / f"{L['prefix']}{i:02d}.png", dest, L["title"], i, total, phrase, is_title, pos, stamps.get(i, []),
                          L.get("title_bottom", 6.42), L.get("title_lines"),
                          L.get("title_max_right"))
        composed.append(dest)
    if only is not None or Ctx.stamp_debug:
        print(f"Lesson {n}: composed slides only (tuning mode), no PPTX written")
        return None
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    for p in composed:
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(str(p), Emu(0), Emu(0), width=SLIDE_W, height=SLIDE_H)
    out = folder / L["pptx"]
    try:
        backup_if_exists(out)
        prs.save(str(out))
    except PermissionError as e:
        print("LOCKED (not updated):", out, "-", e)
        LOCKED.append(str(out))
        out = folder / f"{out.stem}.new-{Ctx.stamp_ts}{out.suffix}"
        prs.save(str(out))   # keep the fresh deck next to the locked one so the copies below still update
    print("wrote", out, "slides:", len(Presentation(str(out)).slides))
    if not copies:
        return out
    if safe_copy(out, folder / f"RemoEd {L['pptx']}"):
        print("lesson folder copy", folder / f"RemoEd {L['pptx']}")
    if safe_copy(out, refs / L["pptx"]):
        print("root copy", refs / L["pptx"])
    stage = refs / "drive-upload-stage"
    stage.mkdir(exist_ok=True)
    if safe_copy(out, stage / f"RemoEd {L['pptx']}"):
        print("drive stage", stage / f"RemoEd {L['pptx']}")
    level_dir = laptop_root / f"Level {LEVEL_NUM}"
    if level_dir.is_dir():
        if safe_copy(out, level_dir / f"RemoEd {L['pptx']}"):
            print("laptop copy", level_dir / f"RemoEd {L['pptx']}")
    else:
        print("laptop folder missing, skipped (not created):", level_dir)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--lessons", nargs="+", type=int, default=sorted(LESSONS))
    ap.add_argument("--root", default=None, help="RemoEdPH project root (default: two levels above this script, else laptop path)")
    ap.add_argument("--font", default=None)
    ap.add_argument("--fallback-font", default=None, help="font for glyphs Arial Rounded lacks (e.g. arrows)")
    ap.add_argument("--laptop-root", default=str(DEFAULT_LAPTOP))
    ap.add_argument("--no-copies", action="store_true")
    ap.add_argument("--stamp-debug", action="store_true", help="outline stamp boxes in red; composes only, no PPTX")
    ap.add_argument("--slides", nargs="+", type=int, default=None, help="compose only these slides (tuning), no PPTX")
    a = ap.parse_args()

    here = Path(__file__).resolve()
    root = Path(a.root) if a.root else (here.parents[2] if (here.parents[2] / "docs" / "lesson-references").is_dir() else DEFAULT_ROOT)
    refs = root / "docs" / "lesson-references"
    Ctx.font_path = first_existing([a.font, r"C:\Windows\Fonts\ARLRDBD.TTF", r"C:\Windows\Fonts\arialbd.ttf"])
    Ctx.fallback_path = first_existing([a.fallback_font, r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\seguisym.ttf",
                                        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"])
    Ctx.logo = refs / "remoedph-logo.jpg"
    Ctx.stamp_debug = a.stamp_debug
    print("root:", root, "| font:", Ctx.font_path, "| fallback:", Ctx.fallback_path)
    if not Ctx.logo.exists():
        raise SystemExit(f"logo not found: {Ctx.logo}")
    for n in a.lessons:
        if n not in LESSONS:
            raise SystemExit(f"Lesson {n} is not configured in this script")
        build(n, refs, Path(a.laptop_root), not a.no_copies, set(a.slides) if a.slides else None)
    if LOCKED:
        print("NOT UPDATED (file locked - close it and re-run):")
        for p in LOCKED:
            print("  ", p)
        raise SystemExit(3)


if __name__ == "__main__":
    main()
