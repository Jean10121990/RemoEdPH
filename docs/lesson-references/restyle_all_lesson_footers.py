"""Restyle bottom-right footer on ALL local RemoEd lesson PPTX files.

Removes the white footer card. Draws SPROUTS! MONTH 1 in white with a drop
shadow, plus the RemoEd PH logo (white plate knocked out) with a drop shadow.
Applies to Sample Lesson and L2M1 Lessons 7–12.
"""
from __future__ import annotations

import io
import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu, Inches

# Reuse chrome helpers from the lesson builder
sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_l2m1_lessons_7_9 import (  # noqa: E402
    LOGO,
    OUT_ROOT,
    PX_H,
    PX_W,
    ROOT,
    SLIDE_H,
    SLIDE_W,
    compose_slide,
    draw_footer,
)

EMU = 914400


def inches(v: int) -> float:
    return v / EMU


SAMPLE_PHRASES = [
    ("Hello!", True),
    ("Happy", False),
    ("Smile", False),
    ("I am happy.", False),
    ("Make a smile!", False),
    ("Happy heart.", False),
    ("Share the joy.", False),
    ("Smile at friends.", False),
    ("Thank you, God.", False),
    ("Touch the happy face!", False),
    ("Touch the happy one!", False),
    ("Your turn!", False),
    ("I am happy.", False),
    ("Happy. Smile.", False),
    ("Goodbye!", False),
]


def extract_bg(slide) -> bytes | None:
    for sh in slide.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            # Prefer full-bleed background (near 0,0 and roughly slide-sized)
            if inches(sh.left) < 0.2 and inches(sh.top) < 0.2 and inches(sh.width) > 10:
                return sh.image.blob
    # Fallback: first picture
    for sh in slide.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return sh.image.blob
    return None


def rebuild_pptx_from_jpgs(jpg_paths: list[Path], out_pptx: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for p in jpg_paths:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(p), Emu(0), Emu(0), width=SLIDE_W, height=SLIDE_H)
    out_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_pptx))
    print("wrote", out_pptx)


def restyle_sample():
    from PIL import Image

    art_dir = OUT_ROOT / "Sample-Lesson" / "art"
    composed_dir = OUT_ROOT / "Sample-Lesson" / "composed"
    extracted = OUT_ROOT / "Sample-Lesson" / "extracted-images"
    art_dir.mkdir(parents=True, exist_ok=True)
    composed_dir.mkdir(parents=True, exist_ok=True)

    title = "Lesson 6 – I am Happy!"
    jpgs = []
    for i in range(1, 16):
        # Prefer clean extracted backgrounds (no chrome baked in)
        clean = extracted / f"slide-{i:02d}-pic-1.jpg"
        art = art_dir / f"sample-slide-{i:02d}.png"
        if clean.exists():
            Image.open(clean).convert("RGB").save(art, "PNG")
            print(f"sample {i}: using extracted bg")
        else:
            src = OUT_ROOT / "Sample-Lesson" / "Sample-Lesson.pptx"
            prs = Presentation(str(src))
            slide = list(prs.slides)[i - 1]
            blob = extract_bg(slide)
            if not blob:
                raise RuntimeError(f"Sample slide {i}: no background picture")
            Image.open(io.BytesIO(blob)).convert("RGB").save(art, "PNG")
            print(f"sample {i}: extracted from pptx")
        phrase, is_title = SAMPLE_PHRASES[i - 1]
        dest = composed_dir / f"sample-slide-{i:02d}.jpg"
        compose_slide(art, dest, title, i, 15, phrase, is_title)
        jpgs.append(dest)

    out = OUT_ROOT / "Sample-Lesson" / "Sample-Lesson.pptx"
    rebuild_pptx_from_jpgs(jpgs, out)
    shutil.copy2(out, ROOT / "docs" / "lesson-references" / "Sample-Lesson.pptx")
    shutil.copy2(jpgs[0], OUT_ROOT / "Sample-Lesson" / "preview-01-title.jpg")
    shutil.copy2(jpgs[1], OUT_ROOT / "Sample-Lesson" / "preview-02-happy.jpg")
    shutil.copy2(jpgs[14], OUT_ROOT / "Sample-Lesson" / "preview-15-garden.jpg")


def restyle_l2m1_lessons():
    from build_l2m1_lessons_7_9 import LESSONS as L79
    from build_l2m1_lessons_10_12 import LESSONS as L1012

    for lesson in list(L79) + list(L1012):
        folder = OUT_ROOT / lesson["folder"]
        art_dir = folder / "art"
        composed_dir = folder / "composed"
        if not art_dir.exists():
            print("skip (no art):", lesson["folder"])
            continue
        jpgs = []
        for i, (phrase, is_title) in enumerate(lesson["pages"], 1):
            src = art_dir / f"{lesson['prefix']}{i:02d}.png"
            if not src.exists():
                # try assets fallback
                assets = Path(
                    r"C:\Users\Window11\.cursor\projects\d-Users-Window11-Documents-RemoEd-Jean-RemoEdPH\assets"
                )
                alt = assets / f"{lesson['prefix']}{i:02d}.png"
                if alt.exists():
                    shutil.copy2(alt, src)
                else:
                    raise FileNotFoundError(src)
            dest = composed_dir / f"{lesson['prefix']}{i:02d}.jpg"
            compose_slide(src, dest, lesson["title"], i, 15, phrase, is_title)
            jpgs.append(dest)
        out = folder / lesson["pptx"]
        rebuild_pptx_from_jpgs(jpgs, out)
        shutil.copy2(out, ROOT / "docs" / "lesson-references" / lesson["pptx"])


def sync_drive_stage():
    stage = ROOT / "docs" / "lesson-references" / "drive-upload-stage"
    stage.mkdir(parents=True, exist_ok=True)
    mapping = {
        "L2M1-Lesson-7/L2M1-Lesson-7-Letter-B-is-for-Brave.pptx": "RemoEd L2M1-Lesson-7-Letter-B-is-for-Brave.pptx",
        "L2M1-Lesson-8/L2M1-Lesson-8-Walking-for-Energy.pptx": "RemoEd L2M1-Lesson-8-Walking-for-Energy.pptx",
        "L2M1-Lesson-9/L2M1-Lesson-9-Good-Morning-Sun.pptx": "RemoEd L2M1-Lesson-9-Good-Morning-Sun.pptx",
        "L2M1-Lesson-10/L2M1-Lesson-10-The-Brave-Bird-Review.pptx": "RemoEd L2M1-Lesson-10-The-Brave-Bird-Review.pptx",
        "L2M1-Lesson-11/L2M1-Lesson-11-I-am-Kind.pptx": "RemoEd L2M1-Lesson-11-I-am-Kind.pptx",
        "L2M1-Lesson-12/L2M1-Lesson-12-Letter-C-is-for-Creative.pptx": "RemoEd L2M1-Lesson-12-Letter-C-is-for-Creative.pptx",
        "Sample-Lesson/Sample-Lesson.pptx": "RemoEd L2M1-Lesson-6-I-am-Happy-Sample.pptx",
    }
    for rel, name in mapping.items():
        src = OUT_ROOT / rel
        if src.exists():
            shutil.copy2(src, stage / name)
            print("staged", name)


def main():
    if not LOGO.exists():
        raise FileNotFoundError(LOGO)
    restyle_sample()
    restyle_l2m1_lessons()
    sync_drive_stage()
    print("done")


if __name__ == "__main__":
    main()
