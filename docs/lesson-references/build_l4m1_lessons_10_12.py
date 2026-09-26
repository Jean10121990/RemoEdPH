r"""Build L4M1 Young Stewards (Month 1 - The Map of My Heart) Lessons 10-12 as RemoEd PPTX decks.

Engine copied from build_l4m1_lessons_07_09.py (from build_l4m1_lessons_04_06.py) (from build_l4m1_lessons_01_03.py) (itself from build_l3m1_lessons_16_17.py) (same chrome as the laptop's shared composer
build_l2m1_lessons_7_9.compose_slide): white lesson pill top-left (0.25",0.2") #1A5696, indigo #5A67D8
"n / N" page badge top-right, yellow #FAD648 learner chip (#1E3A5F), white "YOUNG STEWARDS! MONTH 1"
footer label + knocked-out RemoEd PH logo with drop shadows and NO card (slide-chrome.md), 1920x1080
composed JPG placed full-bleed on a 13.333"x7.5" slide. Level 4 specifics:
  * footer YOUNG STEWARDS! MONTH 1 (codebase level name "Young Stewards"; README/PDF say "Young Steward");
  * page totals follow the brief per lesson: L10 = 20 slides, L11 = 18, L12 = 18;
  * lesson folders lessons/L4M1-Lesson-<N>/ (README pattern L{level}M1-Lesson-{N}, no zero padding),
    art files art/l4m1-l<NN>-slide-<XX>.png, deck "L4M1-Lesson-<N>-<Title>.pptx" + "RemoEd ..." copies,
    laptop copy into ...\A Lesson and Training Materials\Level 4\ (only if that folder exists);
  * title slide: "Lesson N - Title" as stacked white pills low-left at 52-58pt (104-116px), greeting chip top;
  * extra stamp kind "x" (gentle red X over a junk-food item) and chip-less slides (phrase None) where the
    on-screen sentence is stamped in a speech bubble instead;
  * all words are stamped in code; nothing is baked into the art; every overwrite is backed up first.

Saves (three-place pattern): lesson folder lessons/L4M1-Lesson-N/ (plain + "RemoEd " copy), Desktop
"...\A Lesson and Training Materials\Level 4\RemoEd <deck>", and --downloads (default C:\Users\Window11\Downloads)
"RemoEd <deck>"; plus the usual root + drive-upload-stage copies. On the box, --box-out /workspace/remoed/out copies the deck
there and --preview-root /workspace/remoed renders every slide to L4M1-Lesson-N/preview/l4m1-lNN-preview-XX.jpg
(LibreOffice -> PDF -> pdftoppm when available, else the composed JPGs are copied).

Stamp positions can be tuned without editing this file: put a stamps.json next to the art folder.

Usage (laptop, from docs/lesson-references):
    python build_l4m1_lessons_10_12.py --lessons 10 11 12
    python build_l4m1_lessons_10_12.py --lessons 10 --stamp-debug --no-copies   (tuning pass)
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Emu, Inches

FOOTER = "YOUNG STEWARDS! MONTH 1"
LEVEL_NUM = 4
DEFAULT_ROOT = Path(r"D:\Users\Window11\Documents\RemoEd-Jean\RemoEdPH")
DEFAULT_LAPTOP = Path(r"D:\Users\Window11\Desktop\JeanDesktop\RemoEdPH\A Lesson and Training Materials")

NAVY = (26, 86, 150)
NAVY_DARK = (30, 58, 95)
INDIGO = (90, 103, 216)
YELLOW = (250, 214, 72)
WHITE = (255, 255, 255)
PLAQUE_INK = (255, 240, 205)
PLAQUE_HALO = (45, 25, 10)

# (on-screen text, is_title, chip position) - chip: "bottom" | "top" | (centre-x frac, top-y frac)
# On-screen text is verbatim from the brief; quotation marks dropped to match laptop L3M1 Lessons 1-6.
# "title_lines" (optional) sets how the title-slide pills break, for titles too long for "Lesson N -" / "Title".


def _row(words, centres, y0, y1, hw):
    return [{"kind": "word", "text": w, "box": (c - hw, y0, c + hw, y1), "group": "row"} for w, c in zip(words, centres)]


RED = (214, 38, 44)
T = "\u2013"
ORANGE = (238, 118, 28)
CLOCK_GREEN = (120, 245, 150)


def _boxes(words, centres, y0, y1, hw, group="row"):
    return [{"kind": "word", "text": w, "box": (c - hw, y0, c + hw, y1), "group": group} for w, c in zip(words, centres)]


PINK = (226, 72, 150)

LESSONS = {
    # Stamp boxes measured on the FINAL art (flood-fill seeds + gridded crops on the 1280x720 PNGs), 2026-09-26.
    10: {
        "title": f"Lesson 10 {T} Reading Short I CVC Words",
        "title_lines": [f"Lesson 10 {T}", "Reading Short I", "CVC Words"],
        "title_bottom": 7.05,
        "title_max_right": 0.54,
        "folder": "L4M1-Lesson-10",
        "pptx": "L4M1-Lesson-10-Reading-Short-I-CVC-Words.pptx",
        "prefix": "l4m1-l10-slide-",
        "pages": [
            ("Short I CVC Words!", True, (0.24, 0.12)),
            ("Short /i/ Sound: /i/ /i/ /i/", False, (0.28, 0.62)),   # under the disc, left of Remo
            ("p - i - g \u2192 pig", False, "bottom"),
            ("b - i - g \u2192 big", False, "bottom"),
            ("d - i - g \u2192 dig", False, "bottom"),
            ("The big pig can dig.", False, "bottom"),
            ("b - i - n \u2192 bin", False, (0.66, 0.80)),           # right of the bin
            ("Put trash in the bin!", False, (0.18, 0.64)),          # open path left of Sofie
            ("Put trash in the big bin.", False, (0.72, 0.14)),      # sky/palms, clear of Ed and the bin
            ("w - i - g \u2192 wig", False, "bottom"),
            ("Pig | Big | Dig | Bin | Wig", False, "bottom"),
            ("The big pig is near the bin.", False, "bottom"),
            ("God made farm animals.", False, (0.72, 0.13)),         # sky right of Ed
            ("Which word is P - I - G?", False, "top"),              # word boxes sit low (76-89% down)
            ("Action: Pretend to DIG!", False, (0.21, 0.14)),        # clear of Ed, shovel and hands
            ("Where is your waste bin?", False, (0.70, 0.16)),       # keep the pointed-at basket visible
            ("Trace: d - i - g", False, "bottom"),
            ("Short /i/ Review:\npig, big, dig, bin, wig", False, "bottom"),
            ("Eco-Drop &\nDig Badge Earned!", False, (0.22, 0.72)),   # two lines, left of the golden badge (57-70% x)
            ("Keep Earth clean! Goodbye!", False, (0.5, 0.095)),
        ],
        "stamps": {
            2: [{"kind": "word", "text": "I", "box": (0.24, 0.20, 0.36, 0.47), "color": PINK}],                      # pink disc 19-41% x 14-53% (centre 30/33.5%)
            11: [{"kind": "word", "text": w, "box": (x0 + 0.015, 0.676, x1 - 0.015, 0.752), "group": "strips"}         # white strips at the card feet, 66-76.5% down
                 for w, (x0, x1) in zip(["Pig", "Big", "Dig", "Bin", "Wig"],
                                        [(0.061, 0.205), (0.244, 0.389), (0.427, 0.570), (0.608, 0.753), (0.793, 0.943)])],
            14: [{"kind": "word", "text": w, "box": (x0 + 0.02, 0.780, x1 - 0.02, 0.870), "group": "boxes"}           # cream fronts 76-89% down; text stays above the footer row
                 for w, (x0, x1) in zip(["PIG", "PEG", "PAG"], [(0.107, 0.270), (0.416, 0.580), (0.727, 0.895)])],
            17: [{"kind": "dotted", "text": "dig", "box": (0.245, 0.40, 0.54, 0.56), "size": 175, "baseline": 0.549}],   # lines 24-54% x: top ~43%, dashed ~49%, base ~55%
        },
    },
    11: {
        "title": f"Lesson 11 {T} Mapping My Neighborhood",
        "title_lines": [f"Lesson 11 {T}", "Mapping My", "Neighborhood"],
        "title_bottom": 7.05,
        "title_max_right": 0.54,
        "folder": "L4M1-Lesson-11",
        "pptx": "L4M1-Lesson-11-Mapping-My-Neighborhood.pptx",
        "prefix": "l4m1-l11-slide-",
        "pages": [
            ("Mapping My Neighborhood!", True, (0.26, 0.12)),
            ("A neighborhood is where we live and play.", False, "bottom"),
            ("My house is near the park.", False, "bottom"),
            ("House / Home", False, "bottom"),
            ("Park", False, "bottom"),
            ("Store", False, "bottom"),
            ("Street / Road", False, "bottom"),
            ("The store is near my house.", False, "bottom"),
            ("Keep our streets clean!", False, (0.22, 0.14)),        # clear of Ed's hands, wrapper and bag
            ("Be a kind neighbor.", False, "bottom"),
            ("Drive to the Park!", False, "bottom"),
            ("Build the sentence!", False, "top"),                 # brief: Assemble the sentence! (swapped: easier word); tiles reach 81% down
            ("Look left and right\nbefore crossing!", False, (0.68, 0.14)),
            ("What is near\nyour house?", False, (0.80, 0.30)),     # right of Sofie (map board + hands stay clear)
            ("March through\nthe neighborhood!", False, (0.20, 0.14)),
            ("My house is near the park.", False, "bottom"),
            ("Neighborhood Steward Badge!", False, "bottom"),
            ("Be a great neighbor today! Goodbye!", False, (0.5, 0.095)),
        ],
        "stamps": {
            11: [{"kind": "word", "text": "House", "box": (0.240, 0.362, 0.377, 0.417), "group": "map"},            # strip above the house 22.7-39% x 35.1-42.8%
                 {"kind": "word", "text": "Park", "box": (0.668, 0.548, 0.806, 0.603), "group": "map"}],           # right strip 65.5-81.9% x 53.7-61.4% (below the park)
            12: _boxes(["the\nstore.", "My\nhouse", "is\nnear"], [0.270, 0.5025, 0.7345], 0.550, 0.750, 0.064, "tiles"),   # tall tiles 19.5-34.5 / 42.8-57.7 / 66-80.9% x, 48.2-81.2% y
            16: [{"kind": "word", "text": "House", "box": (0.290, 0.295, 0.376, 0.350), "group": "labels"},         # strip 27.8-38.8% x 28.3-36.2%
                 {"kind": "word", "text": "Street", "box": (0.447, 0.456, 0.551, 0.514), "group": "labels"},       # strip 43.5-55.8% x 44-53%
                 {"kind": "word", "text": "Park", "box": (0.796, 0.132, 0.887, 0.190), "group": "labels"},         # strip 78.4-89.9% x 11.9-20.3%
                 {"kind": "word", "text": "Store", "box": (0.819, 0.545, 0.911, 0.608), "group": "labels"}],       # strip 80.7-92.3% x 53.1-62.2%
        },
    },
    12: {
        "title": f"Lesson 12 {T} Planning a Home Eco-Project",
        "title_lines": [f"Lesson 12 {T}", "Planning a", "Home", "Eco-Project"],   # four short pills stay left of Ed (43% x)
        "title_bottom": 7.05,
        "title_max_right": 0.42,
        "folder": "L4M1-Lesson-12",
        "pptx": "L4M1-Lesson-12-Planning-a-Home-Eco-Project.pptx",
        "prefix": "l4m1-l12-slide-",
        "pages": [
            ("Hello!", True, (0.16, 0.12)),
            ("Stretch High!", False, (0.78, 0.62)),
            ("Green Corner", False, (0.22, 0.62)),
            ("Leaf", False, "bottom"),
            ("Pot", False, "bottom"),
            ("Plant", False, "bottom"),
            ("We plan a\ngreen corner.", False, (0.80, 0.62)),   # right of Ed
            ("Dig, Dig!", False, "bottom"),
            ("Water it!", False, (0.80, 0.16)),
            ("God made nature.", False, "bottom"),
            ("Eco-Plan", False, (0.80, 0.70)),
            ("Where is the leaf?", False, "bottom"),
            ("Where is the pot?", False, "bottom"),
            ("Green Leaf\nHunt!", False, (0.125, 0.60)),
            ("Match it!", False, "bottom"),
            ("Leaf. Pot. Plant.", False, "bottom"),
            ("Plan Saved!", False, (0.18, 0.20)),                  # brief: Project Logged! (swapped: easier words)
            ("Goodbye!", False, "bottom"),
        ],
        "stamps": {
            15: [{"kind": "word", "text": t, "box": (0.45, y0, 0.82, y1), "group": "match"}                      # strips 41.9-84.9% x, crossed:
                 for t, (y0, y1) in zip(["Pot", "Plant", "Leaf"], [(0.249, 0.322), (0.438, 0.515), (0.633, 0.707)])],   # rows 24.4-32.7 / 43.3-52 / 62.8-71.2%
        },
    },
}

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
PX_W, PX_H = 1920, 1080


def ix(inches: float) -> int:
    return int(round(inches * PX_W / 13.333))


def iy(inches: float) -> int:
    return int(round(inches * PX_H / 7.5))


class Ctx:
    downloads: Path | None = None
    box_out: Path | None = None
    preview_root: Path | None = None
    font_path: str | None = None
    fallback_path: str | None = None
    logo: Path | None = None
    stamp_debug = False
    stamp_ts = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")


def first_existing(paths):
    for p in paths:
        if p and Path(p).exists():
            return str(p)
    return None


def load_font(size: int, fallback: bool = False):
    path = Ctx.fallback_path if fallback else Ctx.font_path
    return ImageFont.truetype(path, size) if path else ImageFont.load_default()


_MISSING: dict = {}


def has_glyph(font, ch: str) -> bool:
    key = (id(font), ch)
    if key not in _MISSING:
        notdef = font.getmask("\ue000")
        m = font.getmask(ch)
        _MISSING[key] = not (m.size == notdef.size and bytes(m) == bytes(notdef))
    return _MISSING[key]


def runs(text: str, size: int):
    """Split text into (segment, font) runs, using the fallback font for glyphs the main font lacks (e.g. U+2192)."""
    main = load_font(size)
    out, cur, cur_fb = [], "", False
    for ch in text:
        fb = bool(Ctx.fallback_path) and not has_glyph(main, ch)
        if cur and fb != cur_fb:
            out.append((cur, load_font(size, cur_fb)))
            cur = ""
        cur += ch
        cur_fb = fb
    if cur:
        out.append((cur, load_font(size, cur_fb)))
    return out


def measure(draw, text: str, size: int):
    w, top, bot = 0, 10 ** 6, -10 ** 6
    for seg, f in runs(text, size):
        b = draw.textbbox((w, 0), seg, font=f)
        w = b[2]
        top, bot = min(top, b[1]), max(bot, b[3])
    return w, bot - top, top


def draw_runs(draw, xy, text: str, size: int, fill):
    x, y = xy
    for seg, f in runs(text, size):
        draw.text((x, y), seg, font=f, fill=fill)
        x = draw.textbbox((x, y), seg, font=f)[2]


def text_size(draw, text, font):
    b = draw.textbbox((0, 0), text, font=font)
    return b[2] - b[0], b[3] - b[1]


def round_rect(draw, xy, radius, fill):
    draw.rounded_rectangle(xy, radius=radius, fill=fill)


# ---------------------------------------------------------------- safe file writes
def backup_if_exists(dest: Path):
    if dest.exists():
        bdir = dest.parent / "_backup"
        bdir.mkdir(exist_ok=True)
        b = bdir / f"{dest.stem}.{Ctx.stamp_ts}{dest.suffix}"
        shutil.copy2(dest, b)
        print("backup", b)


LOCKED: list = []


def safe_copy(src: Path, dest: Path) -> bool:
    """Copy with backup; a file locked by PowerPoint etc. is reported and skipped, not fatal."""
    try:
        backup_if_exists(dest)
        shutil.copy2(src, dest)
        return True
    except PermissionError as e:
        print("LOCKED (not updated):", dest, "-", e)
        LOCKED.append(str(dest))
        return False


# ---------------------------------------------------------------- footer / chrome
def knock_out_white(im):
    im = im.convert("RGBA")
    px = im.load()
    w, h = im.size
    for y in range(h):
        for x in range(w):
            r, g, b, a = px[x, y]
            if r > 246 and g > 246 and b > 246:
                px[x, y] = (r, g, b, 0)
            elif r > 228 and g > 228 and b > 228:
                px[x, y] = (r, g, b, max(0, a - int((min(r, g, b) - 228) / 18 * 255)))
    return im


_LOGO: dict = {}


def footer_geometry(draw):
    tw, th = text_size(draw, FOOTER, load_font(34))
    logo_w, logo_h = ix(0.78), iy(0.78)
    logo_x = PX_W - ix(0.28) - logo_w
    logo_y = PX_H - iy(0.22) - logo_h
    text_x = logo_x - ix(0.14) - tw
    return tw, th, logo_w, logo_h, logo_x, logo_y, text_x


def draw_footer(img):
    canvas = img.convert("RGBA")
    tw, th, logo_w, logo_h, logo_x, logo_y, text_x = footer_geometry(ImageDraw.Draw(canvas))
    text_y = logo_y + (logo_h - th) // 2 - 2
    font = load_font(34)
    shadow = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    ImageDraw.Draw(shadow).text((text_x, text_y + 2), FOOTER, font=font, fill=(0, 0, 0, 179))  # rgba(0,0,0,0.7)
    canvas = Image.alpha_composite(canvas, shadow.filter(ImageFilter.GaussianBlur(4)))
    sharp = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    ImageDraw.Draw(sharp).text((text_x, text_y), FOOTER, font=font, fill=WHITE + (255,))
    canvas = Image.alpha_composite(canvas, sharp)
    if "logo" not in _LOGO:
        _LOGO["logo"] = knock_out_white(Image.open(Ctx.logo)).resize((logo_w, logo_h), Image.Resampling.LANCZOS)
    logo = _LOGO["logo"]
    ls = Image.new("RGBA", logo.size, (0, 0, 0, 0))
    ls.putalpha(logo.split()[-1].point(lambda a: int(a * 0.6)))  # rgba(0,0,0,0.6)
    pad = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    pad.paste(ls.filter(ImageFilter.GaussianBlur(3)), (logo_x, logo_y + 2), ls.filter(ImageFilter.GaussianBlur(3)))
    canvas = Image.alpha_composite(canvas, pad)
    canvas.paste(logo, (logo_x, logo_y), logo)
    return canvas.convert("RGB")


# ---------------------------------------------------------------- code-stamped text
def fit_size(draw, lines, w, h, start=400, floor=18):
    size = start
    while size > floor:
        f = load_font(size)
        lw = max(text_size(draw, ln, f)[0] for ln in lines)
        lh = int(size * 1.08) * len(lines)
        if lw <= w and lh <= h:
            return size
        size -= 4
    return floor


def stamp(img, s):
    draw = ImageDraw.Draw(img)
    kind = s["kind"]
    if kind == "badge":
        cx, cy = int(PX_W * s["center"][0]), int(PX_H * s["center"][1])
        r = int(PX_W * s.get("r", 0.04))
        sh = Image.new("RGBA", img.size, (0, 0, 0, 0))
        ImageDraw.Draw(sh).ellipse((cx - r, cy - r + 4, cx + r, cy + r + 4), fill=(0, 0, 0, 110))
        img = Image.alpha_composite(img.convert("RGBA"), sh.filter(ImageFilter.GaussianBlur(5))).convert("RGB")
        draw = ImageDraw.Draw(img)
        draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill=YELLOW, outline=WHITE, width=6)
        f = load_font(int(r * 1.2))
        b = draw.textbbox((0, 0), s["text"], font=f)
        draw.text((cx - (b[0] + b[2]) // 2, cy - (b[1] + b[3]) // 2), s["text"], font=f, fill=NAVY_DARK)
        return img
    x0, y0, x1, y1 = (int(PX_W * s["box"][0]), int(PX_H * s["box"][1]), int(PX_W * s["box"][2]), int(PX_H * s["box"][3]))
    if kind == "clone":   # cover a tiny art remnant with the plain background from dy below (soft-edged)
        dy = int(PX_H * s["dy"])
        patch = img.crop((x0, y0 + dy, x1, y1 + dy))
        mask = Image.new("L", patch.size, 0)
        ml, mr = (0 if x0 <= 0 else 6), (0 if x1 >= PX_W else 6)   # no feather on the slide edge itself
        ImageDraw.Draw(mask).rectangle((ml - 8 if ml == 0 else ml, 6, patch.size[0] - mr + (8 if mr == 0 else -1), patch.size[1] - 7), fill=255)
        mask = mask.filter(ImageFilter.GaussianBlur(3))
        if ml == 0 or mr == 0:   # re-solidify the edge column the blur softened
            md = ImageDraw.Draw(mask)
            if ml == 0:
                md.rectangle((0, 8, 3, patch.size[1] - 9), fill=255)
            if mr == 0:
                md.rectangle((patch.size[0] - 4, 8, patch.size[0], patch.size[1] - 9), fill=255)
        img = img.copy()
        img.paste(patch, (x0, y0), mask)
        return img
    if Ctx.stamp_debug:
        draw.rectangle((x0, y0, x1, y1), outline=(255, 0, 0), width=3)
    if kind == "x":   # gentle red X (drawn in code, never in the art) over a not-healthy choice
        w = max(16, int(min(x1 - x0, y1 - y0) * 0.13))
        ends = [((x0, y0), (x1, y1)), ((x1, y0), (x0, y1))]
        base = img.convert("RGBA")
        for col, width, blur, dy in (((0, 0, 0, 90), w + 12, 6, 5), ((255, 255, 255, 235), w + 12, 0, 0), ((220, 45, 50, 240), w, 0, 0)):
            lay = Image.new("RGBA", img.size, (0, 0, 0, 0))
            ld = ImageDraw.Draw(lay)
            for (ax, ay), (bx, by) in ends:
                ld.line((ax, ay + dy, bx, by + dy), fill=col, width=width)
                for px_, py_ in ((ax, ay + dy), (bx, by + dy)):
                    ld.ellipse((px_ - width // 2, py_ - width // 2, px_ + width // 2, py_ + width // 2), fill=col)
            if blur:
                lay = lay.filter(ImageFilter.GaussianBlur(blur))
            base = Image.alpha_composite(base, lay)
        return base.convert("RGB")
    if kind == "lines":
        n = s.get("n", 3)
        gap = (x1 - x0) * 0.12
        seg = ((x1 - x0) - gap * (n - 1)) / n
        th = max(12, (y1 - y0) // 3)
        for k in range(n):
            lx = int(x0 + k * (seg + gap))
            round_rect(draw, (lx - 2, y1 - th - 2, int(lx + seg) + 2, y1 + 2), th // 2 + 2, WHITE)
            round_rect(draw, (lx, y1 - th, int(lx + seg), y1), th // 2, NAVY_DARK)
        return img
    if s.get("plate"):   # clean white rounded label strip drawn in code (art has none)
        rad = (y1 - y0) // 3
        sh = Image.new("RGBA", img.size, (0, 0, 0, 0))
        round_rect(ImageDraw.Draw(sh), (x0, y0 + 6, x1, y1 + 6), rad, (0, 0, 0, 90))
        img = Image.alpha_composite(img.convert("RGBA"), sh.filter(ImageFilter.GaussianBlur(6))).convert("RGB")
        draw = ImageDraw.Draw(img)
        round_rect(draw, (x0, y0, x1, y1), rad, WHITE)
        pad_x, pad_y = int((x1 - x0) * 0.12), int((y1 - y0) * 0.16)
        x0, y0, x1, y1 = x0 + pad_x, y0 + pad_y, x1 - pad_x, y1 - pad_y
    lines = s["text"].split("\n")
    size = s.get("size") or fit_size(draw, lines, x1 - x0, y1 - y0)
    f = load_font(size)
    while s.get("size") and size > 24 and max(draw.textbbox((0, 0), ln, font=f)[2] - draw.textbbox((0, 0), ln, font=f)[0] for ln in lines) > (x1 - x0):
        size -= 4   # fixed sizes are tuned on Liberation; a wider font (Arial Rounded) shrinks to stay inside the box
        f = load_font(size)
    lh = int(size * 1.08)
    total_h = lh * len(lines)
    layer = Image.new("L", img.size, 0)
    ld = ImageDraw.Draw(layer)
    for k, ln in enumerate(lines):
        if s.get("baseline") is not None:   # sit the letters on a drawn guide line (tracing boards)
            b = ld.textbbox((0, 0), ln, font=f, anchor="ls")
            tx = x0 + ((x1 - x0) - (b[2] - b[0])) // 2 - b[0]
            ld.text((tx, int(PX_H * s["baseline"]) + k * lh), ln, font=f, fill=255, anchor="ls")
            continue
        b = ld.textbbox((0, 0), ln, font=f)
        tx = (x0 - b[0]) if s.get("align") == "left" else x0 + ((x1 - x0) - (b[2] - b[0])) // 2 - b[0]
        ty = y0 + ((y1 - y0) - total_h) // 2 + k * lh - b[1] + (lh - (b[3] - b[1])) // 2
        ld.text((tx, ty), ln, font=f, fill=255)
    base = img.convert("RGBA")
    if kind == "glow":
        glow = Image.new("RGBA", img.size, (255, 236, 140, 0))
        glow.putalpha(layer.filter(ImageFilter.MaxFilter(9)).filter(ImageFilter.GaussianBlur(14)))
        base = Image.alpha_composite(base, glow)
        base = Image.alpha_composite(base, glow)
        core = Image.new("RGBA", img.size, (255, 250, 220, 0))
        core.putalpha(layer)
        base = Image.alpha_composite(base, core)
    elif kind == "dotted":
        # tracing style: faint letter fill + dashed outline ring
        from PIL import ImageChops
        k = max(7, (size // 18) | 1)
        ring = ImageChops.subtract(layer.filter(ImageFilter.MaxFilter(k)), layer.filter(ImageFilter.MinFilter(k)))
        step = max(12, size // 9)
        dots = Image.new("L", img.size, 0)
        dd = ImageDraw.Draw(dots)
        for yy in range(y0 - step, y1 + step, step):
            for xx in range(x0 - step, x1 + step, step):
                dd.ellipse((xx, yy, xx + step * 0.62, yy + step * 0.62), fill=255)
        fill = Image.new("RGBA", img.size, (160, 170, 240, 0))
        fill.putalpha(layer.point(lambda v: int(v * 0.35)))
        base = Image.alpha_composite(base, fill)
        col = Image.new("RGBA", img.size, INDIGO + (0,))
        col.putalpha(ImageChops.multiply(ring, dots))
        base = Image.alpha_composite(base, col)
    else:  # word: navy text with a soft white halo so it reads on boards, cards and paper
        halo = Image.new("RGBA", img.size, tuple(s.get("halo", WHITE)) + (0,))
        halo.putalpha(layer.filter(ImageFilter.MaxFilter(5)).point(lambda a: int(a * 0.55)))
        base = Image.alpha_composite(base, halo)
        txt = Image.new("RGBA", img.size, tuple(s.get("color", NAVY_DARK)) + (0,))
        txt.putalpha(layer)
        base = Image.alpha_composite(base, txt)
    return base.convert("RGB")


def uniform_sizes(img, stamps):
    """Stamps sharing a "group" key get the same (largest common) font size, so a row of cards reads evenly."""
    groups: dict = {}
    draw = ImageDraw.Draw(img)
    for st in stamps:
        if st.get("group") and "box" in st and not st.get("size"):
            x0, y0, x1, y1 = (int(PX_W * st["box"][0]), int(PX_H * st["box"][1]), int(PX_W * st["box"][2]), int(PX_H * st["box"][3]))
            sz = fit_size(draw, st["text"].split("\n"), x1 - x0, y1 - y0)
            groups[st["group"]] = min(groups.get(st["group"], 10 ** 6), sz)
    return [dict(st, size=groups[st["group"]]) if st.get("group") in groups else st for st in stamps]


# ---------------------------------------------------------------- slide composition
def load_art(src: Path, zoom=None):
    img = Image.open(src).convert("RGB")
    if zoom:   # (scale, anchor-x frac, anchor-y frac): crop in around a fixed point to push edge decoration off-frame
        sc, ax, ay = zoom
        w, h = img.size
        nw, nh = w / sc, h / sc
        l, t = ax * (w - nw), ay * (h - nh)
        img = img.crop((int(round(l)), int(round(t)), int(round(l + nw)), int(round(t + nh))))
    w, h = img.size
    t = PX_W / PX_H
    if abs(w / h - t) > 0.01:
        if w / h > t:
            nw = int(h * t)
            img = img.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
        else:
            nh = int(w / t)
            img = img.crop((0, (h - nh) // 2, w, (h - nh) // 2 + nh))
    return img.resize((PX_W, PX_H), Image.Resampling.LANCZOS)


def compose_slide(src, dest, title, page, total, phrase, is_title, chip_pos, stamps, title_bottom=6.42, title_lines=None,
                  title_max_right=None, zoom=None, title_top=None):
    img = load_art(src, zoom)
    stamps = uniform_sizes(img, stamps)
    for s in stamps:
        img = stamp(img, s)
    draw = ImageDraw.Draw(img)
    font14 = load_font(36)

    tw, th = text_size(draw, title, font14)
    lw = max(ix(4.2), tw + 44)
    lx, ly = ix(0.25), iy(0.2)
    round_rect(draw, (lx, ly, lx + lw, ly + iy(0.45)), 22, WHITE)
    draw.text((lx + 18, ly + (iy(0.45) - th) // 2 - 2), title, font=font14, fill=NAVY)

    badge = f"{page} / {total}"
    bw, bh, bx, by = ix(1.5), iy(0.45), ix(11.55), iy(0.2)
    round_rect(draw, (bx, by, bx + bw, by + bh), 22, INDIGO)
    tw, th = text_size(draw, badge, font14)
    draw.text((bx + (bw - tw) // 2, by + (bh - th) // 2 - 2), badge, font=font14, fill=WHITE)

    img = draw_footer(img)
    draw = ImageDraw.Draw(img)
    f_tw, _, _, _, _, logo_y, footer_left = footer_geometry(draw)

    if is_title:
        # two stacked pills low-left: "Lesson N -" / "{Title}", 52-58pt (104-116px)
        lines = list(title_lines) if title_lines else [t.strip() for t in title.replace("\u2013", "\u2013\n", 1).split("\n")]
        max_right = footer_left - ix(0.2)
        if title_max_right:   # keep the pills left of a face that sits low in the frame (e.g. Remo)
            max_right = min(max_right, int(PX_W * title_max_right))
        size = 116
        while True:
            f = load_font(size)
            widths = [draw.textbbox((0, 0), ln, font=f)[2] for ln in lines]
            if ix(0.3) + max(widths) + 70 <= max_right or size <= 104:
                break
            size -= 4
        asc, desc = f.getmetrics()
        pill_h, gap = asc + desc + 24, 12
        bottom = iy(title_bottom)
        right = ix(0.3) + max(widths) + 70
        if bottom > logo_y - 12 and right > footer_left - ix(0.15):
            bottom = min(bottom, logo_y - 12)   # never let a low title collide with the footer
        y = bottom - len(lines) * pill_h - (len(lines) - 1) * gap
        if title_top is not None:   # top-anchored stack when a character stands in the lower-left title zone
            y = iy(title_top)
        for ln, w in zip(lines, widths):
            round_rect(draw, (ix(0.3), y, ix(0.3) + w + 70, y + pill_h), 28, WHITE)
            draw.text((ix(0.3) + 35, y + 12), ln, font=f, fill=NAVY)
            y += pill_h + gap

    if phrase:   # None = no learner chip (the sentence is stamped on the art, e.g. a speech bubble)
        size = 64
        chip_lines = phrase.split("\n")        # a chip may wrap onto 2 lines to fit a gap beside a character
        ms = [measure(draw, ln, size) for ln in chip_lines]
        lh, lgap = max(m[1] for m in ms), 10
        tw, th, top = max(m[0] for m in ms), lh * len(ms) + lgap * (len(ms) - 1), min(m[2] for m in ms)
        cw, ch = tw + 72, max(iy(0.7), th + 32)
        if isinstance(chip_pos, tuple):
            cx, cy = int(PX_W * chip_pos[0]) - cw // 2, int(PX_H * chip_pos[1])
        elif chip_pos == "top":
            cx, cy = (PX_W - cw) // 2, iy(0.9)
        else:
            cx, cy = (PX_W - cw) // 2, iy(6.05)
            limit = footer_left - ix(0.15)
            if cx + cw > limit:   # Level 4 footer is wider: slide the chip left on the same bottom row first
                cx = max(ix(0.3), limit - cw)
            if cx + cw > limit:   # still too wide: lift it above the footer row, centred again
                cx, cy = (PX_W - cw) // 2, min(cy, logo_y - 18 - ch)
        round_rect(draw, (cx, cy, cx + cw, cy + ch), 26, YELLOW)
        ly0 = cy + (ch - th) // 2 - top
        for ln, m in zip(chip_lines, ms):
            draw_runs(draw, (cx + (cw - m[0]) // 2, ly0), ln, size, NAVY_DARK)
            ly0 += lh + lgap

    dest.parent.mkdir(parents=True, exist_ok=True)
    backup_if_exists(dest)
    img.save(dest, "JPEG", quality=92)


def build(n: int, refs: Path, laptop_root: Path, copies: bool, only=None):
    L = LESSONS[n]
    folder = refs / "lessons" / L["folder"]
    art_dir, comp_dir = folder / "art", folder / "composed"
    total = len(L["pages"])
    stamps = {int(k): v for k, v in L.get("stamps", {}).items()}
    override = folder / "stamps.json"
    if override.exists():
        for k, v in json.loads(override.read_text(encoding="utf-8")).items():
            stamps[int(k)] = [dict(s, box=tuple(s["box"])) if "box" in s else dict(s, center=tuple(s["center"])) for s in v]
        print("stamp overrides from", override)
    missing = [i for i in range(1, total + 1) if not (art_dir / f"{L['prefix']}{i:02d}.png").exists()]
    if missing:
        raise SystemExit(f"Lesson {n}: missing art for slides {missing} in {art_dir}")
    composed = []
    for i, (phrase, is_title, pos) in enumerate(L["pages"], 1):
        dest = comp_dir / f"{L['prefix']}{i:02d}.jpg"
        if only is None or i in only:
            compose_slide(art_dir / f"{L['prefix']}{i:02d}.png", dest, L["title"], i, total, phrase, is_title, pos, stamps.get(i, []),
                          L.get("title_bottom", 6.42), L.get("title_lines"),
                          L.get("title_max_right"), L.get("art_zoom", {}).get(i), L.get("title_top"))
        composed.append(dest)
    if only is not None or Ctx.stamp_debug:
        print(f"Lesson {n}: composed slides only (tuning mode), no PPTX written")
        return None
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    for p in composed:
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(str(p), Emu(0), Emu(0), width=SLIDE_W, height=SLIDE_H)
    out = folder / L["pptx"]
    try:
        backup_if_exists(out)
        prs.save(str(out))
    except PermissionError as e:
        print("LOCKED (not updated):", out, "-", e)
        LOCKED.append(str(out))
        out = folder / f"{out.stem}.new-{Ctx.stamp_ts}{out.suffix}"
        prs.save(str(out))   # keep the fresh deck next to the locked one so the copies below still update
    print("wrote", out, "slides:", len(Presentation(str(out)).slides))
    if not copies:
        return out
    if safe_copy(out, folder / f"RemoEd {L['pptx']}"):
        print("lesson folder copy", folder / f"RemoEd {L['pptx']}")
    if safe_copy(out, refs / L["pptx"]):
        print("root copy", refs / L["pptx"])
    stage = refs / "drive-upload-stage"
    stage.mkdir(exist_ok=True)
    if safe_copy(out, stage / f"RemoEd {L['pptx']}"):
        print("drive stage", stage / f"RemoEd {L['pptx']}")
    level_dir = laptop_root / f"Level {LEVEL_NUM}"
    if level_dir.is_dir():
        if safe_copy(out, level_dir / f"RemoEd {L['pptx']}"):
            print("laptop copy", level_dir / f"RemoEd {L['pptx']}")
    else:
        print("laptop folder missing, skipped (not created):", level_dir)
    if Ctx.downloads and Ctx.downloads.is_dir():
        if safe_copy(out, Ctx.downloads / f"RemoEd {L['pptx']}"):
            print("downloads copy", Ctx.downloads / f"RemoEd {L['pptx']}")
    if Ctx.box_out:
        Ctx.box_out.mkdir(parents=True, exist_ok=True)
        if safe_copy(out, Ctx.box_out / f"RemoEd {L['pptx']}"):
            print("box copy", Ctx.box_out / f"RemoEd {L['pptx']}")
    if Ctx.preview_root:
        render_previews(n, out, composed)
    return out


def render_previews(n: int, deck: Path, composed):
    """Every slide as l4m1-lNN-preview-XX.jpg: rendered from the finished PPTX (LibreOffice + pdftoppm) when possible."""
    import subprocess
    import tempfile
    pdir = Ctx.preview_root / LESSONS[n]["folder"] / "preview"
    pdir.mkdir(parents=True, exist_ok=True)
    stem = f"l4m1-l{n:02d}-preview-"
    soffice, pdftoppm = shutil.which("soffice") or shutil.which("libreoffice"), shutil.which("pdftoppm")
    if soffice and pdftoppm:
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td) / f"L{n}.pptx"
            shutil.copy2(deck, tmp)
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", td, str(tmp)], check=True,
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            subprocess.run([pdftoppm, "-jpeg", "-r", "96", str(Path(td) / f"L{n}.pdf"), str(Path(td) / "p")], check=True)
            pages = sorted(Path(td).glob("p-*.jpg"), key=lambda p: int(p.stem.split("-")[-1]))
            for i, p in enumerate(pages, 1):
                shutil.copy2(p, pdir / f"{stem}{i:02d}.jpg")
        print("previews rendered:", len(pages), "->", pdir)
    else:
        for i, p in enumerate(composed, 1):
            shutil.copy2(p, pdir / f"{stem}{i:02d}.jpg")
        print("previews copied from composed JPGs:", len(composed), "->", pdir)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--lessons", nargs="+", type=int, default=sorted(LESSONS))
    ap.add_argument("--root", default=None, help="RemoEdPH project root (default: two levels above this script, else laptop path)")
    ap.add_argument("--font", default=None)
    ap.add_argument("--fallback-font", default=None, help="font for glyphs Arial Rounded lacks (e.g. arrows)")
    ap.add_argument("--laptop-root", default=str(DEFAULT_LAPTOP))
    ap.add_argument("--downloads", default=r"C:\Users\Window11\Downloads", help="third save place (skipped if missing)")
    ap.add_argument("--box-out", default=None, help="box only: also copy each deck here (e.g. /workspace/remoed/out)")
    ap.add_argument("--preview-root", default=None, help="box only: render previews into <root>/L4M1-Lesson-N/preview/")
    ap.add_argument("--no-copies", action="store_true")
    ap.add_argument("--stamp-debug", action="store_true", help="outline stamp boxes in red; composes only, no PPTX")
    ap.add_argument("--slides", nargs="+", type=int, default=None, help="compose only these slides (tuning), no PPTX")
    a = ap.parse_args()

    here = Path(__file__).resolve()
    root = Path(a.root) if a.root else (here.parents[2] if (here.parents[2] / "docs" / "lesson-references").is_dir() else DEFAULT_ROOT)
    refs = root / "docs" / "lesson-references"
    Ctx.font_path = first_existing([a.font, r"C:\Windows\Fonts\ARLRDBD.TTF", r"C:\Windows\Fonts\arialbd.ttf"])
    Ctx.fallback_path = first_existing([a.fallback_font, r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\seguisym.ttf",
                                        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"])
    Ctx.logo = refs / "remoedph-logo.jpg"
    Ctx.stamp_debug = a.stamp_debug
    Ctx.downloads = Path(a.downloads) if a.downloads else None
    Ctx.box_out = Path(a.box_out) if a.box_out else None
    Ctx.preview_root = Path(a.preview_root) if a.preview_root else None
    print("root:", root, "| font:", Ctx.font_path, "| fallback:", Ctx.fallback_path)
    if not Ctx.logo.exists():
        raise SystemExit(f"logo not found: {Ctx.logo}")
    for n in a.lessons:
        if n not in LESSONS:
            raise SystemExit(f"Lesson {n} is not configured in this script")
        build(n, refs, Path(a.laptop_root), not a.no_copies, set(a.slides) if a.slides else None)
    if LOCKED:
        print("NOT UPDATED (file locked - close it and re-run):")
        for p in LOCKED:
            print("  ", p)
        raise SystemExit(3)


if __name__ == "__main__":
    main()
