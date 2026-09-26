"""Build L3M1 Saplings Lessons 13-15 as RemoEd PPTX decks.

Generalised from build_l3m1_lesson_12.py. Chrome matches the laptop's shared composer
(build_l2m1_lessons_7_9.compose_slide): white lesson pill (0.25",0.2") 4.2"x0.45" #1A5696,
indigo #5A67D8 page badge (11.55",0.2") 1.5"x0.45", yellow #FAD648 learner chip (#1E3A5F),
white footer label + knocked-out logo with drop shadows and NO card, 1920x1080 composed JPG
placed full-bleed on a 13.333"x7.5" slide. Deliberate differences (same as Lesson 12):
  * title slide: ~52-58pt title as two stacked white pills low-left (not a mid-frame banner
    that covers faces); greeting chip in the open top band;
  * wide body chips are lifted above the footer row instead of colliding with it;
  * art is centre-cropped to 16:9 before resizing (never stretched);
  * phonics / card words are stamped in code (README: never bake text into art);
  * nothing is ever overwritten without a timestamped backup in a sibling "_backup" folder,
    and the D: Level folder is never created (copied only if it exists).

Art: docs/lesson-references/lessons/L3M1-Lesson-<N>/art/l3m1-l<N>-slide-01.png ... -18.png
Stamp positions can be tuned without editing this file: put a stamps.json next to the art
folder ({"4": [{"kind": "lines", ...}], ...}); it replaces that lesson's STAMPS per slide.

Usage (laptop, from docs/lesson-references):
    python build_l3m1_lessons_13_15.py --lessons 13 14 15
    python build_l3m1_lessons_13_15.py --lessons 13 --stamp-debug --no-copies   (tuning pass)
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Emu, Inches

FOOTER = "SAPLINGS! MONTH 1"
LEVEL_NUM = 3
DEFAULT_ROOT = Path(r"D:\Users\Window11\Documents\RemoEd-Jean\RemoEdPH")
DEFAULT_LAPTOP = Path(r"D:\Users\Window11\Desktop\JeanDesktop\RemoEdPH\A Lesson and Training Materials")

NAVY = (26, 86, 150)
NAVY_DARK = (30, 58, 95)
INDIGO = (90, 103, 216)
YELLOW = (250, 214, 72)
WHITE = (255, 255, 255)

# (on-screen text, is_title, chip position) - chip: "bottom" | "top" | (centre-x frac, top-y frac)
# On-screen text is verbatim from the brief; quotation marks dropped to match laptop L3M1 Lessons 1-6.
LESSONS = {
    13: {
        "title": "Lesson 13 \u2013 Writing Short A Words",
        "folder": "L3M1-Lesson-13",
        "pptx": "L3M1-Lesson-13-Writing-Short-A-Words.pptx",
        "prefix": "l3m1-l13-slide-",
        "pages": [
            ("Phonics Writing Time!", True, (0.27, 0.10)),   # left over the sky, clear of the board stamp
            ("Tap Finger Sounds!", False, "bottom"),
            ("Listen \u2192 Say \u2192 Write", False, "bottom"),   # brief: "Listen -> Say -> Write"
            ("/k/ - /a/ - /t/", False, "bottom"),
            ("/m/ - /a/ - /t/", False, "bottom"),
            ("/h/ - /a/ - /t/", False, "bottom"),
            ("Air Write: C - A - T", False, "bottom"),
            ("Write CAT on Paper!", False, "bottom"),
            ("C-A-T, cat.", False, "bottom"),
            ("Do your best always.", False, "bottom"),
            ("Put in order to spell CAT!", False, "bottom"),
            ("/b/ - /a/ - /g/", False, "bottom"),
            ("Show Camera!", False, (0.165, 0.82)),   # beside the paper so both words stay visible
            ("Fast Spelling!", False, "bottom"),
            ("Short A Writer", False, "bottom"),
            ("C-A-T, cat.", False, "bottom"),
            ("CVC Writer Badge Unlocked!", False, "bottom"),
            ("Goodbye!", False, "bottom"),
        ],
        # boxes are fractions of the frame (x0, y0, x1, y1); measured on the final art (2026-09-26)
        "title_bottom": 7.22,   # drop the title pills so line 1 sits below Remo's face (face ~51-63% height)
        "stamps": {
            1: [{"kind": "word", "text": "Aa", "box": (0.60, 0.095, 0.76, 0.235)}],   # clear board strip above Ed/Sofie, left of Grace
            4: [{"kind": "lines", "n": 3, "box": (0.24, 0.655, 0.56, 0.695)}],        # desk top under the kitten
            5: [{"kind": "lines", "n": 3, "box": (0.34, 0.735, 0.66, 0.775)}],        # floor under the mat
            6: [{"kind": "lines", "n": 3, "box": (0.36, 0.66, 0.66, 0.705)}],         # desk under the hat
            7: [{"kind": "glow", "text": "C A T", "box": (0.20, 0.12, 0.55, 0.40)}],  # bright air left of the sparkle trail
            8: [{"kind": "dotted", "text": "CAT", "box": (0.29, 0.28, 0.73, 0.68)}],   # on the paper
            9: [{"kind": "word", "text": "CAT", "box": (0.50, 0.28, 0.79, 0.58)}],     # whiteboard face
            11: [{"kind": "word", "text": "A", "box": (0.262, 0.545, 0.355, 0.715)},  # cube faces
                 {"kind": "word", "text": "T", "box": (0.455, 0.545, 0.550, 0.710)},
                 {"kind": "word", "text": "C", "box": (0.655, 0.535, 0.745, 0.700)}],
            12: [{"kind": "lines", "n": 3, "box": (0.14, 0.645, 0.46, 0.695)}],       # desk left of the bag
            13: [{"kind": "word", "text": "CAT", "box": (0.36, 0.56, 0.68, 0.71)},     # Sofie's paper
                 {"kind": "word", "text": "HAT", "box": (0.36, 0.72, 0.68, 0.87)}],
            14: [{"kind": "word", "text": w, "box": (c - 0.065, 0.50, c + 0.065, 0.66)}
                 for w, c in zip(["CAT", "MAT", "HAT", "BAG"], [0.227, 0.412, 0.594, 0.777])],
            15: [{"kind": "word", "text": w, "box": (c - hw, 0.497, c + hw, 0.545)}    # white label strips under the photos
                 for w, c, hw in zip(["CAT", "MAT", "HAT", "BAG"], [0.160, 0.3075, 0.4325, 0.550], [0.050, 0.045, 0.040, 0.037])],
        },
    },
    14: {
        "title": "Lesson 14 \u2013 Politeness Online & Offline",
        "folder": "L3M1-Lesson-14",
        "pptx": "L3M1-Lesson-14-Politeness-Online-and-Offline.pptx",
        "prefix": "l3m1-l14-slide-",
        "pages": [
            ("Polite Manners Every Day!", True, (0.66, 0.10)),   # right of Ed's head
            ("Big Polite Smile!", False, "bottom"),
            ("Please", False, "bottom"),
            ("Thank You", False, "bottom"),
            ("You\u2019re Welcome", False, "bottom"),
            ("Say please and thank you.", False, "bottom"),
            ("Mute, Listen & Raise Hand!", False, "bottom"),
            ("Manners at Home!", False, "bottom"),
            ("May I have the block, please?", False, "bottom"),
            ("Speak words of kindness.", False, "bottom"),
            ("What do you say?", False, "bottom"),
            ("Thank you, Teacher!", False, "bottom"),
            ("Manners Champion!", False, "bottom"),
            ("Please. Thank You. Welcome.", False, (0.42, 0.82)),   # below the cards (card bottoms ~77%)
            ("Polite Online & Offline", False, "bottom"),
            ("Say please and thank you.", False, "bottom"),
            ("Polite Star Badge Unlocked!", False, "bottom"),
            ("Goodbye!", False, "bottom"),
        ],
        "stamps": {
            11: [{"kind": "badge", "text": "A", "center": (0.058, 0.20), "r": 0.042},   # window/curtain corners, off Ed's head
                 {"kind": "badge", "text": "B", "center": (0.555, 0.20), "r": 0.042}],
            13: [{"kind": "word", "text": "Polite\nManners", "box": (0.33, 0.52, 0.67, 0.76)}],   # wooden sign
            14: [{"kind": "word", "text": w, "box": (c - 0.10, 0.53, c + 0.10, 0.73)}
                 for w, c in zip(["Please", "Thank\nYou", "You\u2019re\nWelcome"], [0.2275, 0.50, 0.7725])],
        },
    },
    15: {
        # vocab confirmed by the user: Different, Good, Together (p.79 header "Friend" ignored);
        # slide 7 keeps Ed/Sofie in their reference look (differences shown via a classmate + Teacher Grace's own glasses)
        "title": "Lesson 15 \u2013 Celebrating Differences",
        "folder": "L3M1-Lesson-15",
        "pptx": "L3M1-Lesson-15-Celebrating-Differences.pptx",
        "prefix": "l3m1-l15-slide-",
        "pages": [
            ("Beautiful Differences!", True, (0.5, 0.10)),
            ("Colorful Crayon Box!", False, "bottom"),
            ("Different", False, "bottom"),
            ("Good", False, "bottom"),
            ("Together", False, "bottom"),
            ("We are all different and good.", False, "bottom"),
            ("Curly, Straight, Glasses!", False, "bottom"),
            ("Fast, Quiet, Creative!", False, "bottom"),
            ("Come play with us!", False, "bottom"),
            ("God made every flower unique.", False, "bottom"),
            ("Which shirt do you like?", False, "bottom"),
            ("Celebrate Your Classmate!", False, "bottom"),
            ("Together as One!", False, "bottom"),
            ("Different. Good. Together.", False, (0.42, 0.82)),   # below the cards (card bottoms ~79%)
            ("Celebrating Differences", False, "bottom"),
            ("We are all different and good.", False, "bottom"),
            ("Rainbow Unity Badge Unlocked!", False, "bottom"),
            ("Goodbye!", False, "bottom"),
        ],
        "stamps": {
            14: [{"kind": "word", "text": w, "box": (c - 0.09, 0.50, c + 0.09, 0.70)}   # three white cards
                 for w, c in zip(["Different", "Good", "Together"], [0.270, 0.5025, 0.7325])],
        },
    },
}

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
PX_W, PX_H = 1920, 1080


def ix(inches: float) -> int:
    return int(round(inches * PX_W / 13.333))


def iy(inches: float) -> int:
    return int(round(inches * PX_H / 7.5))


class Ctx:
    font_path: str | None = None
    fallback_path: str | None = None
    logo: Path | None = None
    stamp_debug = False
    stamp_ts = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")


def first_existing(paths):
    for p in paths:
        if p and Path(p).exists():
            return str(p)
    return None


def load_font(size: int, fallback: bool = False):
    path = Ctx.fallback_path if fallback else Ctx.font_path
    return ImageFont.truetype(path, size) if path else ImageFont.load_default()


_MISSING: dict = {}


def has_glyph(font, ch: str) -> bool:
    key = (id(font), ch)
    if key not in _MISSING:
        notdef = font.getmask("\ue000")
        m = font.getmask(ch)
        _MISSING[key] = not (m.size == notdef.size and bytes(m) == bytes(notdef))
    return _MISSING[key]


def runs(text: str, size: int):
    """Split text into (segment, font) runs, using the fallback font for glyphs the main font lacks (e.g. U+2192)."""
    main = load_font(size)
    out, cur, cur_fb = [], "", False
    for ch in text:
        fb = bool(Ctx.fallback_path) and not has_glyph(main, ch)
        if cur and fb != cur_fb:
            out.append((cur, load_font(size, cur_fb)))
            cur = ""
        cur += ch
        cur_fb = fb
    if cur:
        out.append((cur, load_font(size, cur_fb)))
    return out


def measure(draw, text: str, size: int):
    w, top, bot = 0, 10 ** 6, -10 ** 6
    for seg, f in runs(text, size):
        b = draw.textbbox((w, 0), seg, font=f)
        w = b[2]
        top, bot = min(top, b[1]), max(bot, b[3])
    return w, bot - top, top


def draw_runs(draw, xy, text: str, size: int, fill):
    x, y = xy
    for seg, f in runs(text, size):
        draw.text((x, y), seg, font=f, fill=fill)
        x = draw.textbbox((x, y), seg, font=f)[2]


def text_size(draw, text, font):
    b = draw.textbbox((0, 0), text, font=font)
    return b[2] - b[0], b[3] - b[1]


def round_rect(draw, xy, radius, fill):
    draw.rounded_rectangle(xy, radius=radius, fill=fill)


# ---------------------------------------------------------------- safe file writes
def backup_if_exists(dest: Path):
    if dest.exists():
        bdir = dest.parent / "_backup"
        bdir.mkdir(exist_ok=True)
        b = bdir / f"{dest.stem}.{Ctx.stamp_ts}{dest.suffix}"
        shutil.copy2(dest, b)
        print("backup", b)


LOCKED: list = []


def safe_copy(src: Path, dest: Path) -> bool:
    """Copy with backup; a file locked by PowerPoint etc. is reported and skipped, not fatal."""
    try:
        backup_if_exists(dest)
        shutil.copy2(src, dest)
        return True
    except PermissionError as e:
        print("LOCKED (not updated):", dest, "-", e)
        LOCKED.append(str(dest))
        return False


# ---------------------------------------------------------------- footer / chrome
def knock_out_white(im):
    im = im.convert("RGBA")
    px = im.load()
    w, h = im.size
    for y in range(h):
        for x in range(w):
            r, g, b, a = px[x, y]
            if r > 246 and g > 246 and b > 246:
                px[x, y] = (r, g, b, 0)
            elif r > 228 and g > 228 and b > 228:
                px[x, y] = (r, g, b, max(0, a - int((min(r, g, b) - 228) / 18 * 255)))
    return im


_LOGO: dict = {}


def footer_geometry(draw):
    tw, th = text_size(draw, FOOTER, load_font(34))
    logo_w, logo_h = ix(0.78), iy(0.78)
    logo_x = PX_W - ix(0.28) - logo_w
    logo_y = PX_H - iy(0.22) - logo_h
    text_x = logo_x - ix(0.14) - tw
    return tw, th, logo_w, logo_h, logo_x, logo_y, text_x


def draw_footer(img):
    canvas = img.convert("RGBA")
    tw, th, logo_w, logo_h, logo_x, logo_y, text_x = footer_geometry(ImageDraw.Draw(canvas))
    text_y = logo_y + (logo_h - th) // 2 - 2
    font = load_font(34)
    shadow = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    ImageDraw.Draw(shadow).text((text_x, text_y + 2), FOOTER, font=font, fill=(0, 0, 0, 179))  # rgba(0,0,0,0.7)
    canvas = Image.alpha_composite(canvas, shadow.filter(ImageFilter.GaussianBlur(4)))
    sharp = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    ImageDraw.Draw(sharp).text((text_x, text_y), FOOTER, font=font, fill=WHITE + (255,))
    canvas = Image.alpha_composite(canvas, sharp)
    if "logo" not in _LOGO:
        _LOGO["logo"] = knock_out_white(Image.open(Ctx.logo)).resize((logo_w, logo_h), Image.Resampling.LANCZOS)
    logo = _LOGO["logo"]
    ls = Image.new("RGBA", logo.size, (0, 0, 0, 0))
    ls.putalpha(logo.split()[-1].point(lambda a: int(a * 0.6)))  # rgba(0,0,0,0.6)
    pad = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    pad.paste(ls.filter(ImageFilter.GaussianBlur(3)), (logo_x, logo_y + 2), ls.filter(ImageFilter.GaussianBlur(3)))
    canvas = Image.alpha_composite(canvas, pad)
    canvas.paste(logo, (logo_x, logo_y), logo)
    return canvas.convert("RGB")


# ---------------------------------------------------------------- code-stamped text
def fit_size(draw, lines, w, h, start=400, floor=18):
    size = start
    while size > floor:
        f = load_font(size)
        lw = max(text_size(draw, ln, f)[0] for ln in lines)
        lh = int(size * 1.08) * len(lines)
        if lw <= w and lh <= h:
            return size
        size -= 4
    return floor


def stamp(img, s):
    draw = ImageDraw.Draw(img)
    kind = s["kind"]
    if kind == "badge":
        cx, cy = int(PX_W * s["center"][0]), int(PX_H * s["center"][1])
        r = int(PX_W * s.get("r", 0.04))
        sh = Image.new("RGBA", img.size, (0, 0, 0, 0))
        ImageDraw.Draw(sh).ellipse((cx - r, cy - r + 4, cx + r, cy + r + 4), fill=(0, 0, 0, 110))
        img = Image.alpha_composite(img.convert("RGBA"), sh.filter(ImageFilter.GaussianBlur(5))).convert("RGB")
        draw = ImageDraw.Draw(img)
        draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill=YELLOW, outline=WHITE, width=6)
        f = load_font(int(r * 1.2))
        b = draw.textbbox((0, 0), s["text"], font=f)
        draw.text((cx - (b[0] + b[2]) // 2, cy - (b[1] + b[3]) // 2), s["text"], font=f, fill=NAVY_DARK)
        return img
    x0, y0, x1, y1 = (int(PX_W * s["box"][0]), int(PX_H * s["box"][1]), int(PX_W * s["box"][2]), int(PX_H * s["box"][3]))
    if Ctx.stamp_debug:
        draw.rectangle((x0, y0, x1, y1), outline=(255, 0, 0), width=3)
    if kind == "lines":
        n = s.get("n", 3)
        gap = (x1 - x0) * 0.12
        seg = ((x1 - x0) - gap * (n - 1)) / n
        th = max(12, (y1 - y0) // 3)
        for k in range(n):
            lx = int(x0 + k * (seg + gap))
            round_rect(draw, (lx - 2, y1 - th - 2, int(lx + seg) + 2, y1 + 2), th // 2 + 2, WHITE)
            round_rect(draw, (lx, y1 - th, int(lx + seg), y1), th // 2, NAVY_DARK)
        return img
    lines = s["text"].split("\n")
    size = fit_size(draw, lines, x1 - x0, y1 - y0)
    f = load_font(size)
    lh = int(size * 1.08)
    total_h = lh * len(lines)
    layer = Image.new("L", img.size, 0)
    ld = ImageDraw.Draw(layer)
    for k, ln in enumerate(lines):
        b = ld.textbbox((0, 0), ln, font=f)
        tx = x0 + ((x1 - x0) - (b[2] - b[0])) // 2 - b[0]
        ty = y0 + ((y1 - y0) - total_h) // 2 + k * lh - b[1] + (lh - (b[3] - b[1])) // 2
        ld.text((tx, ty), ln, font=f, fill=255)
    base = img.convert("RGBA")
    if kind == "glow":
        glow = Image.new("RGBA", img.size, (255, 236, 140, 0))
        glow.putalpha(layer.filter(ImageFilter.MaxFilter(9)).filter(ImageFilter.GaussianBlur(14)))
        base = Image.alpha_composite(base, glow)
        base = Image.alpha_composite(base, glow)
        core = Image.new("RGBA", img.size, (255, 250, 220, 0))
        core.putalpha(layer)
        base = Image.alpha_composite(base, core)
    elif kind == "dotted":
        # tracing style: faint letter fill + dashed outline ring
        from PIL import ImageChops
        k = max(7, (size // 18) | 1)
        ring = ImageChops.subtract(layer.filter(ImageFilter.MaxFilter(k)), layer.filter(ImageFilter.MinFilter(k)))
        step = max(12, size // 9)
        dots = Image.new("L", img.size, 0)
        dd = ImageDraw.Draw(dots)
        for yy in range(y0 - step, y1 + step, step):
            for xx in range(x0 - step, x1 + step, step):
                dd.ellipse((xx, yy, xx + step * 0.62, yy + step * 0.62), fill=255)
        fill = Image.new("RGBA", img.size, (160, 170, 240, 0))
        fill.putalpha(layer.point(lambda v: int(v * 0.35)))
        base = Image.alpha_composite(base, fill)
        col = Image.new("RGBA", img.size, INDIGO + (0,))
        col.putalpha(ImageChops.multiply(ring, dots))
        base = Image.alpha_composite(base, col)
    else:  # word: navy text with a soft white halo so it reads on boards, cards and paper
        halo = Image.new("RGBA", img.size, (255, 255, 255, 0))
        halo.putalpha(layer.filter(ImageFilter.MaxFilter(5)).point(lambda a: int(a * 0.55)))
        base = Image.alpha_composite(base, halo)
        txt = Image.new("RGBA", img.size, s.get("color", NAVY_DARK) + (0,))
        txt.putalpha(layer)
        base = Image.alpha_composite(base, txt)
    return base.convert("RGB")


# ---------------------------------------------------------------- slide composition
def load_art(src: Path):
    img = Image.open(src).convert("RGB")
    w, h = img.size
    t = PX_W / PX_H
    if abs(w / h - t) > 0.01:
        if w / h > t:
            nw = int(h * t)
            img = img.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
        else:
            nh = int(w / t)
            img = img.crop((0, (h - nh) // 2, w, (h - nh) // 2 + nh))
    return img.resize((PX_W, PX_H), Image.Resampling.LANCZOS)


def compose_slide(src, dest, title, page, total, phrase, is_title, chip_pos, stamps, title_bottom=6.42):
    img = load_art(src)
    for s in stamps:
        img = stamp(img, s)
    draw = ImageDraw.Draw(img)
    font14 = load_font(36)

    tw, th = text_size(draw, title, font14)
    lw = max(ix(4.2), tw + 44)
    lx, ly = ix(0.25), iy(0.2)
    round_rect(draw, (lx, ly, lx + lw, ly + iy(0.45)), 22, WHITE)
    draw.text((lx + 18, ly + (iy(0.45) - th) // 2 - 2), title, font=font14, fill=NAVY)

    badge = f"{page} / {total}"
    bw, bh, bx, by = ix(1.5), iy(0.45), ix(11.55), iy(0.2)
    round_rect(draw, (bx, by, bx + bw, by + bh), 22, INDIGO)
    tw, th = text_size(draw, badge, font14)
    draw.text((bx + (bw - tw) // 2, by + (bh - th) // 2 - 2), badge, font=font14, fill=WHITE)

    img = draw_footer(img)
    draw = ImageDraw.Draw(img)
    f_tw, _, _, _, _, logo_y, footer_left = footer_geometry(draw)

    if is_title:
        # two stacked pills low-left: "Lesson N -" / "{Title}", 52-58pt (104-116px)
        lines = [t.strip() for t in title.replace("\u2013", "\u2013\n", 1).split("\n")]
        max_right = footer_left - ix(0.2)
        size = 116
        while True:
            f = load_font(size)
            widths = [draw.textbbox((0, 0), ln, font=f)[2] for ln in lines]
            if ix(0.3) + max(widths) + 70 <= max_right or size <= 104:
                break
            size -= 4
        asc, desc = f.getmetrics()
        pill_h, gap = asc + desc + 24, 12
        bottom = iy(title_bottom)
        right = ix(0.3) + max(widths) + 70
        if bottom > logo_y - 12 and right > footer_left - ix(0.15):
            bottom = min(bottom, logo_y - 12)   # never let a low title collide with the footer
        y = bottom - len(lines) * pill_h - (len(lines) - 1) * gap
        for ln, w in zip(lines, widths):
            round_rect(draw, (ix(0.3), y, ix(0.3) + w + 70, y + pill_h), 28, WHITE)
            draw.text((ix(0.3) + 35, y + 12), ln, font=f, fill=NAVY)
            y += pill_h + gap

    size = 64
    tw, th, top = measure(draw, phrase, size)
    cw, ch = tw + 72, max(iy(0.7), th + 32)
    if isinstance(chip_pos, tuple):
        cx, cy = int(PX_W * chip_pos[0]) - cw // 2, int(PX_H * chip_pos[1])
    elif chip_pos == "top":
        cx, cy = (PX_W - cw) // 2, iy(0.9)
    else:
        cx, cy = (PX_W - cw) // 2, iy(6.05)
        if cx + cw > footer_left - ix(0.15):
            cy = min(cy, logo_y - 18 - ch)
    round_rect(draw, (cx, cy, cx + cw, cy + ch), 26, YELLOW)
    draw_runs(draw, (cx + (cw - tw) // 2, cy + (ch - th) // 2 - top), phrase, size, NAVY_DARK)

    dest.parent.mkdir(parents=True, exist_ok=True)
    backup_if_exists(dest)
    img.save(dest, "JPEG", quality=92)


def build(n: int, refs: Path, laptop_root: Path, copies: bool, only=None):
    L = LESSONS[n]
    folder = refs / "lessons" / L["folder"]
    art_dir, comp_dir = folder / "art", folder / "composed"
    total = len(L["pages"])
    stamps = {int(k): v for k, v in L.get("stamps", {}).items()}
    override = folder / "stamps.json"
    if override.exists():
        for k, v in json.loads(override.read_text(encoding="utf-8")).items():
            stamps[int(k)] = [dict(s, box=tuple(s["box"])) if "box" in s else dict(s, center=tuple(s["center"])) for s in v]
        print("stamp overrides from", override)
    missing = [i for i in range(1, total + 1) if not (art_dir / f"{L['prefix']}{i:02d}.png").exists()]
    if missing:
        raise SystemExit(f"Lesson {n}: missing art for slides {missing} in {art_dir}")
    composed = []
    for i, (phrase, is_title, pos) in enumerate(L["pages"], 1):
        dest = comp_dir / f"{L['prefix']}{i:02d}.jpg"
        if only is None or i in only:
            compose_slide(art_dir / f"{L['prefix']}{i:02d}.png", dest, L["title"], i, total, phrase, is_title, pos, stamps.get(i, []),
                          L.get("title_bottom", 6.42))
        composed.append(dest)
    if only is not None or Ctx.stamp_debug:
        print(f"Lesson {n}: composed slides only (tuning mode), no PPTX written")
        return None
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    for p in composed:
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(str(p), Emu(0), Emu(0), width=SLIDE_W, height=SLIDE_H)
    out = folder / L["pptx"]
    try:
        backup_if_exists(out)
        prs.save(str(out))
    except PermissionError as e:
        print("LOCKED (not updated):", out, "-", e)
        LOCKED.append(str(out))
        out = folder / f"{out.stem}.new-{Ctx.stamp_ts}{out.suffix}"
        prs.save(str(out))   # keep the fresh deck next to the locked one so the copies below still update
    print("wrote", out, "slides:", len(Presentation(str(out)).slides))
    if not copies:
        return out
    if safe_copy(out, refs / L["pptx"]):
        print("root copy", refs / L["pptx"])
    stage = refs / "drive-upload-stage"
    stage.mkdir(exist_ok=True)
    if safe_copy(out, stage / f"RemoEd {L['pptx']}"):
        print("drive stage", stage / f"RemoEd {L['pptx']}")
    level_dir = laptop_root / f"Level {LEVEL_NUM}"
    if level_dir.is_dir():
        if safe_copy(out, level_dir / f"RemoEd {L['pptx']}"):
            print("laptop copy", level_dir / f"RemoEd {L['pptx']}")
    else:
        print("laptop folder missing, skipped (not created):", level_dir)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--lessons", nargs="+", type=int, default=sorted(LESSONS))
    ap.add_argument("--root", default=None, help="RemoEdPH project root (default: two levels above this script, else laptop path)")
    ap.add_argument("--font", default=None)
    ap.add_argument("--fallback-font", default=None, help="font for glyphs Arial Rounded lacks (e.g. arrows)")
    ap.add_argument("--laptop-root", default=str(DEFAULT_LAPTOP))
    ap.add_argument("--no-copies", action="store_true")
    ap.add_argument("--stamp-debug", action="store_true", help="outline stamp boxes in red; composes only, no PPTX")
    ap.add_argument("--slides", nargs="+", type=int, default=None, help="compose only these slides (tuning), no PPTX")
    a = ap.parse_args()

    here = Path(__file__).resolve()
    root = Path(a.root) if a.root else (here.parents[2] if (here.parents[2] / "docs" / "lesson-references").is_dir() else DEFAULT_ROOT)
    refs = root / "docs" / "lesson-references"
    Ctx.font_path = first_existing([a.font, r"C:\Windows\Fonts\ARLRDBD.TTF", r"C:\Windows\Fonts\arialbd.ttf"])
    Ctx.fallback_path = first_existing([a.fallback_font, r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\seguisym.ttf",
                                        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"])
    Ctx.logo = refs / "remoedph-logo.jpg"
    Ctx.stamp_debug = a.stamp_debug
    print("root:", root, "| font:", Ctx.font_path, "| fallback:", Ctx.fallback_path)
    if not Ctx.logo.exists():
        raise SystemExit(f"logo not found: {Ctx.logo}")
    for n in a.lessons:
        if n not in LESSONS:
            raise SystemExit(f"Lesson {n} is not configured in this script")
        build(n, refs, Path(a.laptop_root), not a.no_copies, set(a.slides) if a.slides else None)
    if LOCKED:
        print("NOT UPDATED (file locked - close it and re-run):")
        for p in LOCKED:
            print("  ", p)
        raise SystemExit(3)


if __name__ == "__main__":
    main()
